"""Fake workbooks for testing the uploads. Invented names and numbers only."""
from datetime import datetime

import openpyxl

OUT = "/tmp/fake"
import os
os.makedirs(OUT, exist_ok=True)

SITE_HEAD = ["SITE_NAME", "LATITUDE_IN_DECIMAL", "LONGITUDE_IN_DECIMAL", "REGION", "SITE_TYPE"]
DC_HEAD = ["No", "State", "City", "Cluster", "DC NAME / LINK", "LRD", "E2E Status", "Fiber Leg", "lat", "long"]
RES_HEAD = ["Region", "Nodes", "Distance(KM)", "Core Secured", "Expected CV Ready", "Port Info", "Contractor", "Remark",
            "Fiber core monitoring ", "Fiber montiored \nY/N/WIP", "Fiber core requestor", "Fiber core request date",
            "Aging(days)", "No of core requested", "Core assignment", "Core balnce", "Core status", "Loss (dB)", "Tested on"]

SITES = [  # name, lat, lon, region, type
    ("TKAA", 2.9213, 101.6559, "Central", "Data centre"), ("TKAB", 3.1390, 101.6869, "Central", "Data centre"),
    ("TKAC", 3.0733, 101.5185, "Central", "Tower"), ("TKAD", 3.8077, 103.3260, "East", "Tower"),
    ("TKAE", 5.4141, 100.3288, "North", "Data centre"), ("TKAF", 4.5975, 101.0901, "North", "Tower"),
    ("TKAG", 1.4927, 103.7414, "South", "Data centre"), ("TKAH", 2.1896, 102.2501, "South", "Tower"),
    ("TKAI", 31.5, 101.0, "Central", "Tower"),   # problem: coordinates outside Southeast Asia
    ("TKAJ", None, None, None, "Tower"),         # problem: no coordinates, no region
]
DCS = [  # state, city, cluster, name, lrd, e2e, leg
    ("Selangor", "Cyberjaya", "Cyberjaya1", "Cyberjaya DC Alpha", "TKAA", "Completed", "1+1"),
    ("Kuala Lumpur", "Kuala Lumpur", "KL1", "KL DC Beta", "TKAB", "Completed", "1+1"),
    ("Penang", "George Town", "Penang1", "Penang DC Gamma", "TKAE", "In progress", "1+0"),
    ("Johor", "Johor Bahru", "Johor1", "Johor DC Delta", "TKAG", "Completed", "1+1"),
]


def site_master(path, sites, dcs):
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active.append(["Fake site master for upload testing. Not real data."])
    ws = wb.create_sheet("Site Location")
    ws.append(["Site location master (fake)"])
    ws.append(["Source: invented"])
    ws.append([])
    ws.append(SITE_HEAD)  # header on row 4, like the sample structure
    for s in sites:
        ws.append(list(s))
    d = wb.create_sheet("DC List")
    d.append(DC_HEAD)
    for i, x in enumerate(dcs, 1):
        d.append([i, x[0], x[1], x[2], x[3], x[4], x[5], x[6], None, None])
    wb.save(f"{OUT}/{path}")


def reservations(path, rows, title_rows=False, sheet="Reservation"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title_rows:
        ws.append(["Fiber requests report (fake), as at 28 Oct 2026"])
        ws.append([])
    ws.append(RES_HEAD)
    for r in rows:
        ws.append(r)
    wb.save(f"{OUT}/{path}")


def row(nodes, who, count, km=12.5, survey=24, region="Central", assignment=None, date=datetime(2026, 9, 14), status=None):
    return [region, nodes, km, survey, None, "fake port info", "Fake Contractor", "invented", "MON-1", "Y", who, date, 40, count,
            assignment, None, status, None, None]


# 1) site master, 2) a second version to try replace-file behaviour (change, missing, new, possible DC rename)
site_master("fake-site-master-1.xlsx", SITES, DCS)
v2 = [s for s in SITES if s[0] != "TKAH"]                                     # TKAH missing from the file
v2 = [(n, la, lo, "East" if n == "TKAC" else rg, tp) for n, la, lo, rg, tp in v2]  # TKAC region changed
v2.append(("TKAK", 3.4, 102.4, "Central", "Tower"))                            # new site
dcs2 = [(a, b, c, "Cyberjaya DC Alpha 2" if e == "TKAA" else d, e, f, g) for a, b, c, d, e, f, g in DCS]  # same LRD, new name
site_master("fake-site-master-2.xlsx", v2, dcs2)

# 3) reservations: upload AFTER the site master 1 is applied (the sites must exist)
good = [
    row("TKAA-TKAB", "Fake Telco One", 4, assignment="TKAA\nTray 1\nP1-P4\n\nTKAB\nMOD B\nP9-12"),  # core assignment is ignored
    row("TKAB-TKAC", "Fake Telco Two", 2),
    row("TKAC-TKAD", "Fake Cloud Co", 6, region="East"),
    row("TKAE-TKAF", "Fake Telco One", 2, region="North"),
    row("TKAG-TKAH", "Fake Bank Ltd", 8, region="South"),
    row("TKAA-TKAB", "Fake Cloud Co", 2, date=datetime(2026, 10, 2)),   # second customer on the same span
]
bad = [
    row("TKAA-TKAC", "Fake Telco Two", None),        # blank count -> Needs review
    row("TKAA TKAD", "Fake Telco One", 2),           # Nodes not in the form AAAA-BBBB -> Needs review
    row("TKAA-ZZZZ", "Fake Bank Ltd", 2),            # site not in the database -> Needs review
    row("TKAB-TKAD", None, 2),                       # no customer -> Needs review
]
reservations("fake-reservations.xlsx", good + bad)
reservations("fake-reservations-title-row.xlsx", good, title_rows=True, sheet="Fiber requests")  # header on row 3, other sheet name
print(sorted(os.listdir(OUT)))
