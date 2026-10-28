"""TEST-009 / TEST-010: Reservation workbook upload, import switch, re-upload. Synthetic workbooks built here."""
import io
from datetime import datetime

import openpyxl
import pytest
from sqlalchemy import text
from test_auth import mkuser

from app.ingest import reservation_upload as ru

HEAD = ["Region", "Nodes", "Distance(KM)", "Core Secured", "Expected CV Ready", "Port Info", "Contractor", "Remark",
        "Fiber core monitoring ", "Fiber montiored \nY/N/WIP", "Fiber core requestor", "Fiber core request date",
        "Aging(days)", "No of core requested", "Core assignment", "Core balnce", "Core status", "Loss (dB)", "Tested on"]
URL = "/api/admin/uploads"
SWITCH = "/api/admin/reservation-import"
ADMIN = {"X-Dev-User": "boss"}


def ports(a, b, ra="P1-P2", rb="P23-24"):
    return f"{a}\nTray 1\n{ra}\n\n{b}\nMOD B\n{rb}"


def row(nodes="AABB-BBCC", who="Acme", cores="ports", count=None, km=3, survey=4, region="Central", date=datetime(2025, 10, 10),
        status=None, loss=None, tested=None, ra="P1-P2", rb="P23-24"):
    a, b = nodes[:4], nodes[5:9]
    asg = ports(a, b, ra, rb) if cores == "ports" else None
    return [region, nodes, km, survey, None, "raw", "Vendor", "note", "ABC", "Y", who, date, 9, count, asg, 9, status, loss, tested]


def workbook(*rows, extra_sheets=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reservation"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    for name in extra_sheets:
        wb.create_sheet(name).append(["site_name"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
def admin(db):
    mkuser(db, "boss", "admin")
    mkuser(db, "plan", "planner")
    for n in ("AABB", "BBCC", "CCDD"):
        db.execute(text("INSERT INTO sites (name, kind) VALUES (:n, 'site')"), {"n": n})


def up(client, data, name="res.xlsx"):
    return client.post(URL, files={"file": (name, data, "application/octet-stream")}, headers=ADMIN)


def confirm(client, uid, **body):
    return client.post(f"{URL}/{uid}/confirm", json=body, headers=ADMIN)


def rows_of(client, uid, group):
    return client.get(f"{URL}/{uid}/rows", params={"group": group}, headers=ADMIN).json()["items"]


def n(db, sql):
    return db.execute(text(sql)).scalar_one()


def run(client, *rows_, **kw):
    r = up(client, workbook(*rows_, **kw))
    assert r.status_code == 201, r.text
    c = confirm(client, r.json()["upload_id"])
    assert c.status_code == 200, c.text
    return c.json()


def kinds(j):
    return {w["kind"] for w in j["warnings"]}


def test_009_validation_and_nothing_saved(client, db, admin):
    r = up(client, workbook(row(), row(who=" "), row("ZZZZ-AABB"), row(), extra_sheets=("Site Location", "Monitoring")))
    assert r.status_code == 201
    j = r.json()
    uid = j["upload_id"]
    assert j["counts"]["rejected"] == 2 and j["counts"]["reservations"] == 1 and j["counts"]["conflicts"] == 1
    assert j["counts"]["new"] == 1
    assert {"duplicate_port_id", "duplicate_cable", "core_held", "ignored_sheet"} <= kinds(j)
    reasons = {i["row_number"]: i["reason"] for i in rows_of(client, uid, "rejected")}
    assert "blank requestor" in reasons[3] and "unknown site ZZZZ" in reasons[4]
    assert rows_of(client, uid, "rejected")[0]["raw"]["Fiber core requestor"] == " "
    assert client.get(f"{URL}/{uid}", headers=ADMIN).json()["counts"] == j["counts"]
    for t in ("spans", "reservations", "cores", "ports", "customers", "measurements"):
        assert n(db, f"SELECT count(*) FROM {t}") == 0, t
    assert n(db, f"SELECT count(*) FROM upload_rows WHERE upload_id = {uid} AND entity <> 'span'") == 4


def test_unreadable_and_non_xlsx(client, db, admin):
    assert client.post(URL, files={"file": ("a.txt", b"x")}, headers=ADMIN).status_code == 422
    assert up(client, b"not a workbook").status_code == 422
    wb = openpyxl.Workbook()
    out = io.BytesIO()
    wb.save(out)
    assert up(client, out.getvalue()).status_code == 422  # no Reservation sheet


def test_confirm_imports_then_reupload_is_idempotent(client, db, admin):
    out = run(client, row(), row("BBCC-CCDD", who="Beta", cores=None, count=2, survey=4), row("AABB-BBCC (additional core)", ra="P3-P4", rb="P25-26", survey=None, km=None))
    assert out["counts"]["reservations"] == 3 and out["counts"]["spans"] == 2
    assert n(db, "SELECT count(*) FROM reservations WHERE imported") == 3
    assert n(db, "SELECT unassigned_cores FROM reservations WHERE unassigned_span_id IS NOT NULL") == 2
    assert n(db, "SELECT count(DISTINCT user_id) FROM reservation_events") == 1  # the uploading admin
    assert n(db, "SELECT count(*) FROM reservation_events WHERE type IN ('requested','approved')") == 6
    before = {t: n(db, f"SELECT count(*) FROM {t}") for t in ("spans", "ports", "cores", "reservations", "customers", "core_holds")}
    snap = db.execute(text("SELECT * FROM reservations ORDER BY id")).all()
    events = n(db, "SELECT count(*) FROM reservation_events")
    r = up(client, workbook(row(), row("BBCC-CCDD", who="Beta", cores=None, count=2, survey=4), row("AABB-BBCC (additional core)", ra="P3-P4", rb="P25-26", survey=None, km=None)))
    j = r.json()
    assert j["counts"]["already_imported"] == 3 and j["counts"]["reservations"] == 0 and j["counts"]["new"] == 0
    assert confirm(client, j["upload_id"]).json()["counts"]["reservations"] == 0
    assert {t: n(db, f"SELECT count(*) FROM {t}") for t in before} == before
    assert db.execute(text("SELECT * FROM reservations ORDER BY id")).all() == snap
    assert n(db, "SELECT count(*) FROM reservation_events") == events
    assert confirm(client, j["upload_id"]).status_code == 409


def test_missing_rows_untouched_and_protected_fields(client, db, admin):
    run(client, row(), row("BBCC-CCDD", who="Beta", ra="P5-P6", rb="P7-8"))
    sid = n(db, "SELECT s.id FROM spans s JOIN sites a ON a.id = s.site_a_id WHERE a.name = 'AABB'")
    boss = n(db, "SELECT id FROM users WHERE username = 'boss'")
    db.execute(text("UPDATE spans SET distance_km = 9 WHERE id = :i"), {"i": sid})
    db.execute(text("INSERT INTO data_edits (entity, entity_id, field, old_value, new_value, reason, user_id) "
                    "VALUES ('span', :i, 'distance_km', '3', '9', 'surveyed again', :u)"), {"i": sid, "u": boss})
    res_before = db.execute(text("SELECT * FROM reservations ORDER BY id")).all()
    # second file lacks BBCC-CCDD, differs on distance (protected) and region (not protected)
    j = up(client, workbook(row(km=5, region="North"))).json()
    uid = j["upload_id"]
    assert j["counts"]["missing"] == 1 and j["counts"]["changed"] == 1 and j["counts"]["already_imported"] == 1
    ch = rows_of(client, uid, "changed")[0]
    by = {c["field"]: c for c in ch["data"]["changes"]}
    assert by["distance_km"]["protected"] and by["distance_km"]["note"] == "edited in dashboard, file differs"
    assert (by["distance_km"]["old"], by["distance_km"]["new"]) == (9.0, 5.0)
    assert not by["region"]["protected"] and ch["data"]["edited_in_dashboard"]
    assert rows_of(client, uid, "missing")[0]["data"]["key"] == "BBCC-CCDD"
    confirm(client, uid)
    assert n(db, f"SELECT distance_km FROM spans WHERE id = {sid}") == 9
    assert n(db, f"SELECT region FROM spans WHERE id = {sid}") == "North"
    assert n(db, "SELECT reason FROM data_edits WHERE entity = 'span' AND field = 'region'").startswith("reservation upload res.xlsx #")
    assert n(db, "SELECT count(*) FROM spans") == 2 and n(db, "SELECT count(*) FROM reservations") == 2  # missing span untouched
    # now the admin ticks Excel wins
    j = up(client, workbook(row(km=5, region="North"))).json()
    row_id = rows_of(client, j["upload_id"], "changed")[0]["id"]
    bad = confirm(client, j["upload_id"], excel_wins=[{"row_id": row_id, "field": "region"}])
    assert bad.status_code == 422  # region is not protected
    assert confirm(client, j["upload_id"], excel_wins=[{"row_id": row_id, "field": "distance_km"}]).status_code == 200
    assert float(n(db, f"SELECT distance_km FROM spans WHERE id = {sid}")) == 5
    assert n(db, "SELECT count(*) FROM data_edits WHERE entity = 'span' AND field = 'distance_km' AND reason LIKE 'reservation upload%'") == 1
    assert db.execute(text("SELECT * FROM reservations ORDER BY id")).all() == res_before


def test_measurements_appended_once(client, db, admin):
    run(client, row(loss=0.5, tested=datetime(2026, 1, 5)))
    assert n(db, "SELECT count(*) FROM measurements") == 1
    run(client, row(loss=0.5, tested=datetime(2026, 1, 5)))
    assert n(db, "SELECT count(*) FROM measurements") == 1
    run(client, row(loss=0.7, tested=datetime(2026, 3, 1)))
    assert n(db, "SELECT count(*) FROM measurements") == 2
    assert n(db, "SELECT count(*) FROM measurements WHERE upload_id IS NOT NULL") == 2


def test_faulty_row_sends_held_reservation_to_action_needed(client, db, admin):
    run(client, row())
    j = up(client, workbook(row(status="Faulty"))).json()
    assert j["counts"]["faulty"] == 1 and j["counts"]["reservations"] == 0
    assert any("Faulty" in w["message"] and "Action needed" in w["message"] for w in j["warnings"])
    assert "faulty_held" not in [i["kind"] for i in client.get("/api/needs-review", headers=ADMIN).json()]
    out = confirm(client, j["upload_id"]).json()
    assert out["counts"]["faulty_cores"] == 2 and out["counts"]["reservations"] == 0
    items = client.get("/api/needs-review", headers=ADMIN).json()
    assert "faulty_held" in [i["kind"] for i in items]
    assert n(db, "SELECT count(*) FROM reservations") == 1 and n(db, "SELECT count(*) FROM reservation_events WHERE type NOT IN ('requested','approved')") == 0
    again = up(client, workbook(row(status="Faulty"))).json()  # already marked: nothing more to do
    assert again["counts"]["faulty"] == 0 and again["counts"]["already_imported"] == 1


def test_switch_closed(client, db, admin):
    assert client.get(SWITCH, headers=ADMIN).json()["closed"] is False
    assert client.post(f"{SWITCH}/close", json={"reason": " "}, headers=ADMIN).status_code == 422
    j = up(client, workbook(row(), row("BBCC-CCDD", who="Beta"))).json()  # staged while open
    r = client.post(f"{SWITCH}/close", json={"reason": "dashboard is the source now"}, headers=ADMIN)
    assert r.status_code == 200 and r.json()["closed"] and r.json()["closed_by"] == "boss"
    assert client.post(f"{SWITCH}/close", json={"reason": "again"}, headers=ADMIN).status_code == 409
    st = client.get(SWITCH, headers=ADMIN).json()
    assert st["closed"] and st["reason"] == "dashboard is the source now" and st["closed_at"]
    assert n(db, "SELECT count(*) FROM data_edits WHERE entity = 'system' AND entity_id = 0 AND field = 'reservation_import'") == 1
    out = confirm(client, j["upload_id"]).json()  # re-checked at confirm
    assert out["counts"]["reservations"] == 0 and out["ignored_reservation_rows"] == 2
    assert n(db, "SELECT count(*) FROM reservations") == 0 and n(db, "SELECT count(*) FROM spans") == 2  # spans still feed
    j = up(client, workbook(row(), row("BBCC-CCDD", who="Beta"))).json()
    assert j["switch_closed"] and j["counts"]["ignored"] == 2 and j["counts"]["reservations"] == 0
    assert "import_closed" in kinds(j) and len([w for w in j["warnings"] if w["kind"] == "import_closed"]) == 1
    assert rows_of(client, j["upload_id"], "reservations") == []


def test_planner_forbidden(client, db, admin):
    P = {"X-Dev-User": "plan"}
    f = {"file": ("a.xlsx", workbook(row()))}
    assert client.post(URL, files=f, headers=P).status_code == 403
    for method, url in [("get", f"{URL}/1"), ("get", f"{URL}/1/rows?group=new"), ("post", f"{URL}/1/confirm"),
                        ("get", SWITCH), ("post", f"{SWITCH}/close")]:
        assert getattr(client, method)(url, headers=P, **({"json": {"reason": "x"}} if method == "post" else {})).status_code == 403, url
    assert client.get(f"{URL}/999", headers=ADMIN).status_code == 404


def test_count_only_and_capacity(client, db, admin):
    out = run(client, row(cores=None, count=3, survey=4), row(cores=None, count=2, survey=4, who="Beta"))
    assert out["counts"]["reservations"] == 1  # second asks 2 but only 1 free
    j = up(client, workbook(row(cores=None, count=3, survey=4), row(cores=None, count=2, survey=4, who="Beta"))).json()
    assert j["counts"]["already_imported"] == 1 and j["counts"]["needs_review"] == 1
    assert "only 1 free" in rows_of(client, j["upload_id"], "needs_review")[0]["reason"]
    assert n(db, "SELECT unassigned_cores FROM reservations") == 3 and n(db, "SELECT count(*) FROM core_holds") == 0
    assert n(db, "SELECT survey_cores FROM spans") == 4 and n(db, "SELECT count(*) FROM cores") == 4
    j = up(client, workbook(row("BBCC-CCDD", cores=None, count=1, survey=None))).json()
    assert j["counts"]["needs_review"] == 1 and "not surveyed" in rows_of(client, j["upload_id"], "needs_review")[0]["reason"]
    dup = up(client, workbook(row("BBCC-CCDD", cores=None, count=1, survey=5), row("BBCC-CCDD", cores=None, count=1, survey=5))).json()
    assert dup["counts"]["reservations"] == 1 and dup["counts"]["needs_review"] == 1


def test_held_core_conflict_skipped_at_confirm(client, db, admin):
    a = up(client, workbook(row(who="Acme"))).json()["upload_id"]
    b = up(client, workbook(row(who="Beta", ra="P2-P3", rb="P24-25"))).json()["upload_id"]
    assert confirm(client, a).json()["counts"]["reservations"] == 1
    out = confirm(client, b).json()
    assert out["counts"]["reservations"] == 0 and out["skipped"][0]["row_number"] == 2
    assert "already taken" in out["skipped"][0]["reason"]
    assert n(db, "SELECT count(*) FROM reservations") == 1 and n(db, "SELECT count(*) FROM core_holds") == 2


def test_same_cores_different_details_needs_review(client, db, admin):
    run(client, row(who="Acme"))
    j = up(client, workbook(row(who="Beta"))).json()
    assert j["counts"]["conflicts"] == 0 and j["counts"]["needs_review"] == 1
    assert "differ" in rows_of(client, j["upload_id"], "needs_review")[0]["reason"]


def test_confirm_is_atomic(client, db, admin, monkeypatch):
    uid = up(client, workbook(row(), row("BBCC-CCDD", who="Beta"))).json()["upload_id"]
    real, calls = ru.place_reservation, []

    def boom(*a, **k):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("disk on fire")
        return real(*a, **k)

    monkeypatch.setattr(ru, "place_reservation", boom)
    with pytest.raises(RuntimeError):
        confirm(client, uid)
    for t in ("spans", "reservations", "cores", "ports", "customers"):
        assert n(db, f"SELECT count(*) FROM {t}") == 0, t
    assert n(db, f"SELECT status FROM uploads WHERE id = {uid}") == "staged"
    monkeypatch.undo()
    assert confirm(client, uid).status_code == 200


def test_site_master_kind_not_visible(client, db, admin):
    uid = db.execute(text("INSERT INTO uploads (user_id, filename, kind) VALUES ((SELECT id FROM users WHERE username='boss'), 'x', 'site_master') RETURNING id")).scalar_one()
    assert client.get(f"{URL}/{uid}", headers=ADMIN).status_code == 404
    assert confirm(client, uid).status_code == 404


def test_header_row_and_sheet_name_found_by_scanning(client, db, admin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiber requests"  # not called Reservation
    ws.append(["Report as at 2026-10-27"])
    ws.append([])
    ws.append([None, None, *HEAD])  # headers start below row 1 and right of column A
    ws.append([None, None, *row()])
    out = io.BytesIO()
    wb.save(out)
    r = up(client, out.getvalue())
    assert r.status_code == 201, r.text
    assert r.json()["counts"]["reservations"] == 1
