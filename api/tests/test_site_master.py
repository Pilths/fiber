"""TEST-019: site master upload (replace-file model). Fixtures are tiny workbooks built here; no real data."""
import io

import openpyxl
import pytest
from alembic import command
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from app.ingest import site_master as sm
from conftest import alembic_cfg
from test_auth import PW, login, mkuser

SITE_HEAD = ["SITE_NAME", "LATITUDE_IN_DECIMAL", "LONGITUDE_IN_DECIMAL", "REGION", "SITE_TYPE"]
DC_HEAD = ["No", "State", "DC NAME / LINK", "LRD", "lat", "long"]
URL = "/api/admin/site-master"


def site(name, lat=3.1, lon=101.6, region="Central", kind="Tower"):
    return [name, lat, lon, region, kind]


def dc(name, lrd, lat=None, lon=None, state="Selangor"):
    return [1, state, name, lrd, lat, lon]


def workbook(sites=None, dcs=None, site_tab="Site Location", dc_tab="DC List", site_head=SITE_HEAD, dc_head=DC_HEAD,
             site_offset=0, extra_tab=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws.append(["not a data sheet"])
    if sites is not None:
        s = wb.create_sheet(site_tab)
        for _ in range(site_offset):
            s.append(["title row"])
        s.append(site_head)
        for r in sites:
            s.append(r)
    if dcs is not None:
        d = wb.create_sheet(dc_tab)
        d.append(dc_head)
        for r in dcs:
            d.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture
def admin(client, db):
    uid = mkuser(db, "boss", "admin")
    assert login(client, "boss").status_code == 200
    return uid


def post(client, data, name="sm.xlsx"):
    return client.post(f"{URL}/uploads", files={"file": (name, data, "application/octet-stream")})


def sheets(r):
    return {s["kind"]: s for s in r.json()["sheets"]}


def apply(client, uid, **body):
    return client.post(f"{URL}/uploads/{uid}/apply", json=body)


def count(db, sql="SELECT count(*) FROM sites"):
    return db.execute(text(sql)).scalar_one()


def seed_sites(db, names, kind="site"):
    for n in names:
        db.execute(text("INSERT INTO sites (name, kind, geom, details) VALUES (:n, :k, ST_SetSRID(ST_MakePoint(101.6, 3.1), 4326), "
                        "CAST('{\"region\": \"Central\", \"SITE_TYPE\": \"Tower\"}' AS jsonb))"), {"n": n, "k": kind})


def rows(client, uid, sheet, group, **q):
    return client.get(f"{URL}/uploads/{uid}/rows", params={"sheet": sheet, "group": group, **q}).json()


def row_of(client, uid, sheet, group, key):
    return next(i for i in rows(client, uid, sheet, group)["items"] if i["key"] == key)


def test_sheets_found_by_headers_not_name_or_row(client, admin):
    data = workbook([["AAAA", 3.1, 101.6, "Central", "x"]], [[1, "DC One", "AAAA", 3.2, 101.7]], site_tab="zz1", dc_tab="zz2",
                    site_offset=3, site_head=["Site Name", " latitude in  decimal", "LONGITUDE_in_decimal", "Region", "x"],
                    dc_head=["No", "dc name/link", "lrd", "Lat", "Long"])
    s = sheets(post(client, data))
    assert s["sites"]["found"] and s["sites"]["sheet_name"] == "zz1" and s["sites"]["counts"]["added"] == 1
    assert s["dcs"]["found"] and s["dcs"]["sheet_name"] == "zz2" and s["dcs"]["counts"]["added"] == 1


def test_header_row_found_below_row_ten_and_away_from_column_a(client, admin):
    pad = [None, None, None]
    data = workbook([pad + site("AAAA")], None, site_offset=24, site_head=pad + SITE_HEAD)
    s = sheets(post(client, data))
    assert s["sites"]["found"] and s["sites"]["counts"]["added"] == 1


def test_no_matching_sheet_blocks_that_part_with_expected_headers(client, admin):
    s = sheets(post(client, workbook([site("AAAA")], None)))
    assert not s["dcs"]["found"] and s["dcs"]["blocked"] and not s["dcs"]["needs_override"]
    assert "DC Name / Link, LRD" in s["dcs"]["block_reasons"][0]
    s = sheets(post(client, workbook(None, [dc("DC One", "AAAA")])))
    assert "site_name, latitude_in_decimal, longitude_in_decimal" in s["sites"]["block_reasons"][0]
    assert s["sites"]["blocked"] and s["dcs"]["found"]


def test_placeholders_are_blank_not_invalid_text(client, admin):
    data = workbook([["AAAA", "TBA", "N/A", "North", "-"], ["BBBB", 3.1, 101.6, "na", "TBC"], ["CCCC", 3.1, 101.6, "North", "tba"]],
                    [dc("DC One", "CCCC", "TBC", "-")])
    r = post(client, data)
    p = {i["key"]: i for i in rows(client, r.json()["upload_id"], "sites", "problem")["items"]}
    assert p["AAAA"]["reason"] == "missing coordinates"
    assert p["BBBB"]["reason"] == "region is blank"
    ok = row_of(client, r.json()["upload_id"], "sites", "added", "CCCC")
    assert "SITE_TYPE" not in ok["parsed"]["details"]
    assert sheets(r)["dcs"]["counts"]["added"] == 1  # blank DC coordinates fall back to the site


def test_preview_groups_and_nothing_written_before_apply(client, admin, db):
    seed_sites(db, ["KEEP", "OLDD", "SAME"])
    db.execute(text("UPDATE sites SET geom = ST_SetSRID(ST_MakePoint(101.7, 3.2), 4326) WHERE name = 'KEEP'"))
    before = (count(db), count(db, "SELECT count(*) FROM data_edits"))
    data = workbook([site("KEEP"), site("SAME"), site("NEWW"), site("BADD", 0, 0), site("ODDR", region="Atlantis")])
    r = post(client, data)
    uid = r.json()["upload_id"]
    c = sheets(r)["sites"]["counts"]
    assert (c["rows"], c["added"], c["changed"], c["missing"], c["problem"], c["unchanged"]) == (5, 1, 1, 1, 2, 1)
    ch = row_of(client, uid, "sites", "changed", "KEEP")["changes"]
    assert ch == [{"field": "location", "old": "3.200000,101.700000", "new": "3.100000,101.600000", "protected": False}]
    assert row_of(client, uid, "sites", "missing", "OLDD")["db_id"]
    assert "unknown region" in row_of(client, uid, "sites", "problem", "ODDR")["reason"]
    assert row_of(client, uid, "sites", "problem", "BADD")["raw"]["LATITUDE_IN_DECIMAL"] == 0
    assert rows(client, uid, "sites", "problem", page_size=1)["total"] == 2
    assert len(rows(client, uid, "sites", "problem", page_size=1)["items"]) == 1
    assert (count(db), count(db, "SELECT count(*) FROM data_edits")) == before  # staging only
    assert count(db, "SELECT count(*) FROM upload_rows") == 6  # 5 file rows raw + 1 missing


def test_ten_percent_drop_blocks_until_override(client, admin, db):
    seed_sites(db, [f"S{i:03d}" for i in range(10)])
    r = post(client, workbook([site(f"S{i:03d}") for i in range(9)]))
    s = sheets(r)["sites"]
    assert s["blocked"] and s["needs_override"] and "fewer" in s["block_reasons"][0]
    uid = r.json()["upload_id"]
    assert apply(client, uid).status_code == 409
    assert apply(client, uid, override_drop={"sites": True}).status_code == 422  # reason required
    ok = apply(client, uid, override_drop={"sites": True}, reason="S009 was decommissioned")
    assert ok.status_code == 200 and ok.json()["sheets"]["sites"]["deactivated"] == 1
    assert count(db, "SELECT active FROM sites WHERE name = 'S009'") is False
    assert count(db, "SELECT count(*) FROM data_edits WHERE entity = 'upload' AND field = 'override_drop'") == 1
    # 9% fewer is allowed
    seed_sites(db, [f"T{i:03d}" for i in range(11)])
    db.execute(text("UPDATE sites SET active = false WHERE name LIKE 'S%'"))
    r = post(client, workbook([site(f"T{i:03d}") for i in range(10)]))
    assert not sheets(r)["sites"]["blocked"]


@pytest.mark.parametrize("data, why", [
    (workbook([]), "empty"),
    (workbook([site("AAAA"), site("AAAA", 3.5)]), "duplicate"),
])
def test_empty_and_duplicate_keys_block_and_cannot_be_overridden(client, admin, db, data, why):
    r = post(client, data)
    s = sheets(r)["sites"]
    assert s["blocked"] and not s["needs_override"] and why in s["block_reasons"][0]
    assert apply(client, r.json()["upload_id"], override_drop={"sites": True}, reason="x").status_code == 409
    assert count(db) == 0


def test_missing_headers_block_not_overridable(client, admin):
    r = post(client, workbook([site("AAAA")], site_head=["SITE_NAME", "LAT", "LONG", "REGION", "X"]))
    s = sheets(r)["sites"]
    assert not s["found"] and s["blocked"] and not s["needs_override"]
    assert apply(client, r.json()["upload_id"], override_drop={"sites": True}, reason="x").status_code == 409


def test_failing_dc_sheet_does_not_block_good_sites_sheet(client, admin, db):
    r = post(client, workbook([site("AAAA"), site("BBBB")], [dc("DC One", "AAAA"), dc("DC One", "BBBB")]))
    s = sheets(r)
    assert not s["sites"]["blocked"] and s["dcs"]["blocked"]
    ok = apply(client, r.json()["upload_id"])
    assert ok.status_code == 200
    assert ok.json()["sheets"]["sites"]["added"] == 2 and ok.json()["sheets"]["dcs"]["applied"] is False
    assert count(db) == 2


def test_dc_location_falls_back_to_site_and_creates_no_site(client, admin, db):
    seed_sites(db, ["GOOD", "BADS"])
    db.execute(text("UPDATE sites SET geom = NULL WHERE name = 'BADS'"))
    sites_before = count(db, "SELECT count(*) FROM sites WHERE kind = 'site'")
    data = workbook([site("GOOD"), site("BADS", 0, 0)],
                    [dc("DC own", "GOOD", 3.3, 101.8), dc("DC fall", "GOOD", "TBA", None), dc("DC swapped", "GOOD", 101.8, 3.3),
                     dc("DC none", "BADS", None, None), dc("DC ghost", "ZZZZ", 3.3, 101.8)])
    r = post(client, data)
    uid = r.json()["upload_id"]
    assert row_of(client, uid, "dcs", "added", "DC own")["location_note"] is None
    assert row_of(client, uid, "dcs", "added", "DC fall")["location_note"] == "location from site GOOD"
    assert row_of(client, uid, "dcs", "added", "DC swapped")["location_note"] == "location from site GOOD"
    none = row_of(client, uid, "dcs", "problem", "DC none")
    assert none["status"] == "needs_review" and "no valid coordinates" in none["reason"]
    assert "unknown site code ZZZZ" in row_of(client, uid, "dcs", "problem", "DC ghost")["reason"]
    assert apply(client, uid).status_code == 200
    assert count(db, "SELECT count(*) FROM sites WHERE kind = 'site'") == sites_before  # DC rows add no LRD site
    assert count(db, "SELECT count(*) FROM sites WHERE kind = 'data_center'") == 3
    assert count(db, "SELECT lrd_site_id = (SELECT id FROM sites WHERE name = 'GOOD') FROM sites WHERE name = 'DC fall'")
    assert count(db, "SELECT geom IS NULL FROM sites WHERE name = 'DC fall'")  # drawn from the linked site
    assert count(db, "SELECT ST_Y(geom) FROM sites WHERE name = 'DC own'") == pytest.approx(3.3)


def test_site_with_bad_coordinates_is_needs_review_and_not_applied(client, admin, db):
    r = post(client, workbook([site("GOOD"), site("BADD", 0, 0)]))
    assert apply(client, r.json()["upload_id"]).status_code == 200
    assert count(db, "SELECT count(*) FROM sites WHERE name = 'BADD'") == 0


def test_apply_updates_with_history_and_deactivates_missing(client, admin, db):
    seed_sites(db, ["AAAA", "GONE", "BACK"])
    db.execute(text("UPDATE sites SET active = false WHERE name = 'BACK'"))
    data = workbook([site("AAAA", 3.5, 101.9, "North", "Pole"), site("BACK"), site("NEWW")], None)
    r = post(client, data, "master-oct.xlsx")
    uid = r.json()["upload_id"]
    res = apply(client, uid)
    assert res.status_code == 200
    assert res.json()["sheets"]["sites"] | {"applied": True} == {"added": 1, "changed": 2, "protected_kept": 0, "renamed": 0,
                                                                  "deactivated": 1, "kept_active": 0, "skipped": 0, "applied": True}
    assert count(db, "SELECT active FROM sites WHERE name = 'GONE'") is False  # never deleted
    assert count(db, "SELECT active FROM sites WHERE name = 'BACK'") is True
    assert count(db, "SELECT details->>'region' FROM sites WHERE name = 'AAAA'") == "North"
    edits = db.execute(text("SELECT field, old_value, new_value, reason FROM data_edits WHERE entity = 'site' "
                            "AND entity_id = (SELECT id FROM sites WHERE name = 'AAAA') ORDER BY field")).all()
    assert [e[0] for e in edits] == ["details.SITE_TYPE", "location", "region"]
    assert all(e[3] == f"site master upload master-oct.xlsx #{uid}" for e in edits)
    assert db.execute(text("SELECT user_id FROM data_edits WHERE field = 'region'")).scalar_one() == admin
    assert apply(client, uid).status_code == 409  # already applied
    cur = client.get(f"{URL}/current").json()
    assert cur["current"]["filename"] == "master-oct.xlsx" and cur["current"]["sheets"]["sites"]["rows"] == 3
    assert cur["current"]["uploaded_by"] == "boss"


def test_missing_site_in_use_stays_active_with_warning(client, admin, db):
    seed_sites(db, ["AAAA", "BBBB", "CCCC"])
    db.execute(text("INSERT INTO spans (site_a_id, site_b_id) SELECT a.id, b.id FROM sites a, sites b "
                    "WHERE a.name = 'BBBB' AND b.name = 'CCCC'"))
    r = post(client, workbook([site("AAAA")], None))
    uid = r.json()["upload_id"]
    assert row_of(client, uid, "sites", "missing", "BBBB")["in_use"] is True
    res = apply(client, uid, override_drop={"sites": True}, reason="test").json()
    assert res["sheets"]["sites"]["kept_active"] == 2 and res["sheets"]["sites"]["deactivated"] == 0
    assert any("BBBB" in w for w in res["warnings"])
    assert count(db, "SELECT count(*) FROM sites WHERE active") == 3


def test_dashboard_edited_field_is_protected(client, admin, db):
    seed_sites(db, ["AAAA"])
    sid = count(db, "SELECT id FROM sites WHERE name = 'AAAA'")
    db.execute(text("INSERT INTO data_edits (entity, entity_id, field, old_value, new_value, reason, user_id) "
                    "VALUES ('site', :i, 'location', 'a', 'b', 'surveyed on site', :u)"), {"i": sid, "u": admin})
    r = post(client, workbook([site("AAAA", 3.5, 101.9, "North")], None))
    uid = r.json()["upload_id"]
    row = row_of(client, uid, "sites", "changed", "AAAA")
    loc = next(c for c in row["changes"] if c["field"] == "location")
    assert row["edited_in_dashboard"] and loc["protected"] and loc["note"] == "edited in dashboard, file differs"
    res = apply(client, uid).json()["sheets"]["sites"]
    assert res["changed"] == 1 and res["protected_kept"] == 1
    assert count(db, "SELECT ST_Y(geom) FROM sites WHERE name = 'AAAA'") == pytest.approx(3.1)  # dashboard value kept
    assert count(db, "SELECT details->>'region' FROM sites WHERE name = 'AAAA'") == "North"


def rename_setup(client, db):
    seed_sites(db, ["AAAA"])
    seed_sites(db, ["Old DC"], "data_center")
    db.execute(text("UPDATE sites SET lrd_site_id = (SELECT id FROM sites WHERE name = 'AAAA'), "
                    "geom = ST_SetSRID(ST_MakePoint(101.8, 3.3), 4326) WHERE name = 'Old DC'"))
    r = post(client, workbook([site("AAAA")], [dc("New DC", "AAAA", 3.3, 101.8, "Johor")]))
    return r.json()


def test_undecided_possible_rename_is_refused(client, admin, db):
    up = rename_setup(client, db)
    res = apply(client, up["upload_id"], override_drop={"dcs": True}, reason="x")
    assert res.status_code == 422 and "Old DC" in res.json()["detail"] and "New DC" in res.json()["detail"]
    assert count(db, "SELECT count(*) FROM sites WHERE name = 'New DC'") == 0


def test_possible_rename_flagged_and_separate_decision(client, admin, db):
    up = rename_setup(client, db)
    pr = up["possible_renames"]
    assert len(pr) == 1 and pr[0]["missing_dc_name"] == "Old DC" and pr[0]["new_dc_name"] == "New DC" and pr[0]["lrd"] == "AAAA"
    assert row_of(client, up["upload_id"], "dcs", "added", "New DC")["possible_rename"]["missing_dc_id"] == pr[0]["missing_dc_id"]
    res = apply(client, up["upload_id"], renames=[{"missing_dc_id": pr[0]["missing_dc_id"], "new_row_id": pr[0]["new_row_id"],
                                                    "decision": "separate"}], override_drop={"dcs": True}, reason="x")
    assert res.status_code == 200, res.text
    assert count(db, "SELECT active FROM sites WHERE name = 'Old DC'") is False
    assert count(db, "SELECT count(*) FROM sites WHERE name = 'New DC'") == 1


def test_possible_rename_rename_decision_keeps_history(client, admin, db):
    up = rename_setup(client, db)
    pr = up["possible_renames"][0]
    old_id = pr["missing_dc_id"]
    res = apply(client, up["upload_id"], override_drop={"dcs": True}, reason="x",
                renames=[{"missing_dc_id": old_id, "new_row_id": pr["new_row_id"], "decision": "rename"}])
    assert res.status_code == 200, res.text
    assert res.json()["sheets"]["dcs"]["renamed"] == 1
    assert count(db, "SELECT name FROM sites WHERE id = :i".replace(":i", str(old_id))) == "New DC"
    assert count(db, "SELECT active FROM sites WHERE id = :i".replace(":i", str(old_id))) is True
    assert count(db, "SELECT count(*) FROM sites WHERE kind = 'data_center'") == 1
    assert count(db, f"SELECT old_value || '>' || new_value FROM data_edits WHERE entity_id = {old_id} AND field = 'name'") == "Old DC>New DC"
    bad = apply(client, up["upload_id"])
    assert bad.status_code == 409  # already applied
    other = post(client, workbook([site("AAAA")], [dc("New DC", "AAAA", 3.3, 101.8)])).json()
    assert apply(client, other["upload_id"], renames=[{"missing_dc_id": 999999, "new_row_id": 1, "decision": "rename"}],
                 override_drop={"dcs": True}, reason="x").status_code == 422


def test_apply_is_all_or_nothing(client, admin, db, monkeypatch):
    real = sm.apply_sheet

    def boom(db_, admin_id, kind, *a, **k):
        out = real(db_, admin_id, kind, *a, **k)
        if kind == "dcs":
            raise RuntimeError("boom")
        return out

    monkeypatch.setattr(sm, "apply_sheet", boom)
    r = post(client, workbook([site("AAAA")], [dc("DC One", "AAAA", 3.3, 101.8)]))
    uid = r.json()["upload_id"]
    with pytest.raises(RuntimeError):
        apply(client, uid)
    assert count(db) == 0
    assert count(db, f"SELECT status FROM uploads WHERE id = {uid}") == "staged"
    monkeypatch.undo()
    assert apply(client, uid).status_code == 200 and count(db) == 2


def test_planner_gets_403_and_bad_files_are_refused(client, db, admin):
    client.cookies.clear()
    mkuser(db, "pat", "planner")
    login(client, "pat")
    assert post(client, workbook([site("AAAA")])).status_code == 403
    assert client.get(f"{URL}/current").status_code == 403
    assert client.get(f"{URL}/uploads/1/rows", params={"sheet": "sites", "group": "added"}).status_code == 403
    assert apply(client, 1).status_code == 403
    client.cookies.clear()
    login(client, "boss", PW)
    assert post(client, b"not a workbook", "x.xlsx").status_code == 422
    assert post(client, workbook([site("AAAA")]), "x.csv").status_code == 422


def test_migration_0006_roundtrip(engine, db):
    cfg = alembic_cfg()
    command.downgrade(cfg, "0005")
    command.upgrade(cfg, "head")
    u = db.execute(text("INSERT INTO users (username, display_name, role) VALUES ('u','U','admin') RETURNING id")).scalar_one()
    assert db.execute(text("INSERT INTO sites (name, kind) VALUES ('A','site') RETURNING active")).scalar_one() is True
    db.execute(text("INSERT INTO uploads (user_id, filename, kind) VALUES (:u, 'f.xlsx', 'site_master')"), {"u": u})
    c = db.execute(text("INSERT INTO customers (name) VALUES ('X') RETURNING id")).scalar_one()
    db.execute(text("INSERT INTO reservations (customer_id, planner_id) VALUES (:c, :u)"), {"c": c, "u": u})
    with pytest.raises(DBAPIError, match="append-only"):
        db.execute(text("UPDATE reservations SET created_at = now()"))


def test_timing_36k_sites(client, admin, db):
    import time
    sites = [site(f"S{i:05d}", 1.5 + (i % 200) * 0.03, 100.0 + (i // 200) * 0.01) for i in range(36000)]
    dcs = [dc(f"DC {i:03d}", f"S{i:05d}", None, None) for i in range(100)]
    data = workbook(sites, dcs)
    t = time.time()
    r = post(client, data)
    t1 = time.time() - t
    uid = r.json()["upload_id"]
    t = time.time()
    page = rows(client, uid, "sites", "added", page_size=100)
    t2 = time.time() - t
    t = time.time()
    res = apply(client, uid)
    t3 = time.time() - t
    print(f"TIMING upload={t1:.1f}s rows_page={t2:.2f}s apply={t3:.1f}s")
    assert res.status_code == 200 and page["total"] == 36000 and count(db) == 36100
