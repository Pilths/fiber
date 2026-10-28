from datetime import date, datetime

import openpyxl
import pytest
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from app import seed
from app.ingest import checks
from app.ingest.excel import parse_workbook

RES_HEAD = ["Region", "Nodes", "Distance(KM)", "Core Secured", "Expected CV Ready", "Port Info", "Contractor", "Remark",
            "Fiber core monitoring ", "Fiber montiored \nY/N/WIP", "Fiber core requestor", "Fiber core request date",
            "Aging(days)", "No of core requested", "Core assignment", "Core balnce"]
SITE_HEAD = ["SITE_NAME", "LATITUDE_IN_DECIMAL", "LONGITUDE_IN_DECIMAL", "SITE_TYPE", "DESCRIPTION"]
DC_HEAD = ["No", "State", "City", "Cluster", "DC NAME / LINK", "LRD", "E2E Status", "Fiber Leg", "lat", "long"]
PORTS = "{a} \nTray 1\nP1-P2\n\n{b}\nMOD B\nP23-24"


def res(region, nodes, who="Acme", ports=None, **kw):
    ports = ports or PORTS.format(a=nodes[:4], b=nodes[5:9])
    return [region, nodes, kw.get("km", 3), 4, 45940, "raw port info", "Vendor", "note", "ABC", kw.get("fm", "Y"),
            who, datetime(2025, 10, 10), 99, 2, ports, 99]


@pytest.fixture
def workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reservation"
    ws.append(RES_HEAD)
    ws.append([None] * 7 + ["DWDM Trunk Route 1"])  # title row
    ws.append(res("Nothern", "AABB-BBCC"))
    ws.append(res(None, "BBCC-AABB", ports=PORTS.format(a="BBCC", b="AABB")))  # reversed pair, merged region below
    ws.append(res(None, "AABB-BBCC (additional core)", ports="AABB\nP3-P4\n\nBBCC\nP25-26"))
    ws.merge_cells("A3:A5")
    ws.append([None] * 16)
    ws.append(res("Southern", "AABB-BBCC", ports="AABB\nP5-P6\n\nBBCC\nP23-282"))
    ws.append(res("Atlantis", "AABB-BBCC"))
    ws.append(res("East coast", "AABB-BBCC", who="  "))
    ws.append(res("Central", "AABB-ZZZZ"))
    ws.append([None] * 7 + ["Company A project"])
    ws.append(res("Central", "CCDD-AABB", who="Beta", fm="maybe"))
    ws.append(res("Central", "CCDD-AABB", who="Beta", ports="CCDD\nP1-P2\n\nAABB\nP1-P2"))
    sl = wb.create_sheet("Site Location")
    sl.append(["Title row"])
    sl.append([])
    sl.append([])
    sl.merge_cells("A1:C3")
    sl.append(SITE_HEAD)
    for row in [("AABB", "3.1", "101.6", "Tower", "main"), ("BBCC", 3.15, 101.7, None, None), ("GGHH", 101.7, 3.2, None, None),  # swapped
                ("CCDD", 0, 0, None, None), ("EEFF", 4.5, 100.3, None, None), (None, 1, 1, None, None),
                ("EEFF", 4.5, 100.3, None, None)]:
        sl.append(list(row))
    dc = wb.create_sheet("DC List")
    dc.append(DC_HEAD)
    dc.append([1, "Selangor", "Cyberjaya", "C1", "Cyberjaya DC 1", "AABB", "Completed", "1+1", "2.9", "101.65"])
    wb.create_sheet("Monitoring").append(["Nodes", "x"])
    path = tmp_path / "fixture.xlsx"
    wb.save(path)
    return path


def by_row(rows, sheet):
    return {r.row_number: r for r in rows if r.sheet == sheet}


def test_checks():
    assert checks.check_coords(3.1, 101.6) is None
    assert "swapped" in checks.check_coords(101.6, 3.1)
    assert checks.check_coords(0, 0) and checks.check_coords(None, 1) and checks.check_coords(50, 50)
    assert [checks.map_region(v)[0] for v in ("Nothern", " northern", "SOUTHERN", "Souther", "East  coast", "Central")] == \
        ["North", "North", "South", "South", "East Coast", "Central"]
    assert checks.map_region("Atlantis")[0] is None
    assert checks.split_nodes("AABB-BBCC") == ("AABB", "BBCC", False)
    assert checks.split_nodes("aabb-bbcc (Additional CORE)") == ("aabb", "bbcc", True)
    assert checks.split_nodes("AABB") is None and checks.split_nodes(None) is None
    assert checks.duplicate_names(["a", "b", "a"]) == {"a"}
    assert checks.check_known_sites(["A", "B"], {"A"}) and not checks.check_known_sites(["A"], None)


def test_reservation_sheet(workbook):
    rows = by_row(parse_workbook(workbook), "Reservation")
    assert sorted(rows) == [3, 4, 5, 7, 8, 9, 10, 12, 13]  # title and blank rows are not data rows
    r3, r4, r5 = rows[3], rows[4], rows[5]
    assert r3.parsed["region"] == r4.parsed["region"] == r5.parsed["region"] == "North"  # merged + Nothern
    assert r3.parsed["group_label"] == "DWDM Trunk Route 1" and rows[13].parsed["group_label"] == "Company A project"
    assert (r3.status, r3.reason) == ("ready", None)
    assert "ports" not in r3.parsed and r3.parsed["unassigned_cores"] == 2  # Core assignment is kept raw, not parsed
    assert "Tray 1" in r3.raw["Core assignment"]
    assert r3.parsed["cv_ready_date"] == date(2025, 10, 10) and r3.parsed["planned_date"] == date(2025, 10, 10)
    assert r3.parsed["distance_km"] == 3 and r3.parsed["survey_cores"] == 4
    assert r3.parsed["port_info_raw"] == "raw port info" and r3.parsed["contractor"] == "Vendor" and r3.parsed["monitoring_code"] == "ABC"
    assert "Aging(days)" in r3.raw and "Aging(days)" not in r3.parsed and "core balance" not in r3.parsed
    assert r3.raw["Fiber montiored Y/N/WIP"] == "Y" and r3.raw["Fiber core request date"] == "2025-10-10T00:00:00"
    assert (r4.parsed["site_a"], r4.parsed["site_b"]) == ("BBCC", "AABB")
    assert r5.parsed["additional"] is True and r5.parsed["site_b"] == "BBCC" and r5.status == "ready"
    assert rows[7].status == "ready"  # a malformed port range no longer matters: ports are not parsed
    assert rows[8].status == "needs_review" and "Atlantis" in rows[8].reason
    assert rows[9].status == "rejected" and "blank requestor" in rows[9].reason
    assert rows[10].status == "needs_review" and "ZZZZ" in rows[10].reason
    assert rows[12].status == "needs_review" and "fiber monitored" in rows[12].reason


def test_site_and_dc_sheets(workbook):
    rows = parse_workbook(workbook)
    assert {r.sheet for r in rows} == {"Reservation", "Site Location", "DC List"}  # Monitoring skipped
    sites = by_row(rows, "Site Location")
    assert sorted(sites) == [5, 6, 7, 8, 9, 11]  # header on row 4; row 10 has no name
    assert sites[5].status == "ready" and sites[5].parsed["details"] == {"SITE_TYPE": "Tower", "DESCRIPTION": "main"}
    assert sites[7].status == "needs_review" and "swapped" in sites[7].reason
    assert sites[8].status == "needs_review" and "0,0" in sites[8].reason
    assert sites[6].status == sites[9].status == "ready"
    assert sites[11].status == "needs_review" and "duplicate" in sites[11].reason
    dc = by_row(rows, "DC List")[2]
    assert dc.status == "ready" and dc.parsed["kind"] == "data_center" and dc.parsed["lrd"] == "AABB"
    assert dc.parsed["lat"] == 2.9 and dc.parsed["details"]["E2E Status"] == "Completed" and dc.parsed["details"]["Fiber Leg"] == "1+1"


def test_unreadable_cells_and_known_sites(workbook):
    wb = openpyxl.load_workbook(workbook)
    wb["Reservation"]["D3"] = "four"
    wb["Reservation"]["L3"] = "not a date"
    wb.save(workbook)
    r = by_row(parse_workbook(workbook, known_sites={"ZZZZ"}), "Reservation")
    assert r[3].status == "needs_review" and "core secured" in r[3].reason and "request date" in r[3].reason
    assert r[10].status == "ready"  # ZZZZ is known now


def test_row_exception_marks_needs_review(workbook, monkeypatch):
    def boom(v):
        if v == "Atlantis":
            raise RuntimeError("boom")
        return "North", None
    monkeypatch.setattr(checks, "map_region", boom)
    r = by_row(parse_workbook(workbook), "Reservation")
    assert r[8].status == "needs_review" and "boom" in r[8].reason and r[3].status == "ready"


def test_excel_check_reports_and_writes_nothing(workbook, capsys, monkeypatch):
    from app.ingest import excel_check
    monkeypatch.setattr("sys.argv", ["excel_check", str(workbook)])
    excel_check.main()
    out = capsys.readouterr().out
    assert "Reservation: 9 rows" in out and "row 9 [rejected] blank requestor" in out and "row 7 [needs_review] latitude and longitude look swapped" in out
    assert "ready: 5, needs_review: 3, rejected: 1" in out  # reservations


def test_seed_loads_ready_rows(db, workbook):
    counts, skipped = seed.load(db, str(workbook))
    assert counts == {"sites": 3, "data_centers": 1, "spans": 1, "customers": 1, "reservations": 4, "cores": 4, "ports": 0}
    assert len(skipped) == 8  # 3 sites + 4 reservation rows not ready, and row 13 whose site CCDD was not loaded
    assert any("row 13 [skipped]" in s for s in skipped)
    n = lambda sql: db.execute(text(sql)).scalar_one()
    assert n("SELECT count(*) FROM core_holds") == 0 == n("SELECT count(*) FROM reservation_cores")  # all count-only
    assert n("SELECT count(*) FROM reservation_events WHERE type = 'approved'") == 4
    assert n("SELECT s.name FROM sites d JOIN sites s ON s.id = d.lrd_site_id WHERE d.kind = 'data_center'") == "AABB"
    assert n("SELECT details->>'SITE_TYPE' FROM sites WHERE name = 'AABB'") == "Tower"
    assert n("SELECT survey_cores FROM spans") == 4 and n("SELECT region FROM spans") == "North"
    assert n("SELECT count(*) FROM reservations WHERE additional AND group_label = 'DWDM Trunk Route 1'") == 1
    with pytest.raises(DBAPIError, match="append-only"):
        db.execute(text("UPDATE reservations SET remark = 'x'"))


def test_seed_refuses_non_empty(db, workbook):
    db.execute(text("INSERT INTO sites (name, kind) VALUES ('X', 'site')"))
    with pytest.raises(SystemExit):
        seed.load(db, str(workbook))


@pytest.fixture
def count_workbook(workbook):
    """The fixture plus two rows: blank assignment with a count (count-only), and blank assignment with no count."""
    wb = openpyxl.load_workbook(workbook)
    ws = wb["Reservation"]
    count_only = res("Central", "AABB-BBCC", who="Gamma")
    count_only[13], count_only[14] = 3, None
    neither = res("Central", "AABB-BBCC", who="Gamma")
    neither[13], neither[14] = None, None
    ws.append(count_only)
    ws.append(neither)
    wb.save(workbook)
    return workbook


def test_count_is_read_and_blank_count_is_needs_review(count_workbook):
    rows = by_row(parse_workbook(count_workbook), "Reservation")
    assert (rows[14].status, rows[14].reason) == ("ready", None)
    assert rows[14].parsed["unassigned_cores"] == 3 and not rows[14].parsed.get("ports")
    assert rows[15].status == "needs_review" and "No of core requested is blank" in rows[15].reason
    assert rows[15].parsed["unassigned_cores"] == 0


def test_seed_surveyed_cores_and_count_only(db, count_workbook):
    counts, skipped = seed.load(db, str(count_workbook))
    n = lambda sql: db.execute(text(sql)).scalar_one()
    assert counts["reservations"] == 5 and any("row 15 [needs_review]" in s for s in skipped)
    assert n("SELECT survey_cores FROM spans") == 4 == n("SELECT count(*) FILTER (WHERE number <= 4) FROM cores")
    assert n("SELECT count(*) FROM cores") == 4 and n("SELECT count(*) FROM cores WHERE port_a_id IS NOT NULL") == 0
    row = db.execute(text("SELECT unassigned_cores, unassigned_span_id = (SELECT id FROM spans) FROM reservations "
                          "WHERE unassigned_cores = 3")).one()
    assert tuple(row) == (3, True)
    rid = n("SELECT id FROM reservations WHERE unassigned_cores = 3")
    assert db.execute(text("SELECT type FROM reservation_events WHERE reservation_id = :r ORDER BY id"), {"r": rid}).scalars().all() == ["requested", "approved"]
    assert n(f"SELECT count(*) FROM reservation_cores WHERE reservation_id = {rid}") == 0
    assert n(f"SELECT count(*) FROM core_holds WHERE reservation_id = {rid}") == 0
    assert n(f"SELECT remaining FROM reservation_unassigned_v WHERE reservation_id = {rid}") == 3
