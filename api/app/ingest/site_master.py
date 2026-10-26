"""Site master workbook (replace-file model): find sheets by headers, parse, diff against the database, stage, apply.

Staging writes only `uploads` and `upload_rows`. Apply recomputes the diff from the staged rows against the
current database, so a stale preview cannot overwrite newer data.
"""
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any

import openpyxl
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ingest import checks
from app.ingest.checks import blank, norm
from app.ingest.excel import _json
from app.routing.graph import invalidate as invalidate_routes

BLANKS = {"", "tba", "n/a", "na", "-", "tbc"}
HEADER_SCAN_ROWS = 10
DROP_LIMIT = (9, 10)  # a sheet with file rows * 10 <= database rows * 9 (10% or more fewer) is blocked
UPLOAD_REASON = "site master upload"
LOCATION_FIELDS = {"location", "geom", "lat", "lon", "latitude", "longitude"}
SHEETS = {  # kind: (required header keys, expected-headers text, entity, own header keys)
    "sites": (("sitename", "latitudeindecimal", "longitudeindecimal"), "site_name, latitude_in_decimal, longitude_in_decimal",
              "site", {"sitename", "latitudeindecimal", "longitudeindecimal", "region"}),
    "dcs": (("dcnamelink", "lrd"), "DC Name / Link, LRD", "dc",
            {"dcnamelink", "lrd", "lat", "long", "latitude", "longitude", "lon"}),
}
GROUP_COL = {"added": "new", "changed": "changed", "missing": "missing", "problem": "needs_review", "unchanged": "unchanged"}


def isblank(v: Any) -> bool:
    """Empty, TBA, N/A, NA, -, TBC (any case) are blank, never invalid text."""
    return v is None or str(v).strip().casefold() in BLANKS


def hkey(h: Any) -> str:
    return re.sub(r"[^0-9a-z]", "", str(h).casefold())


def _fmt(lat: float | None, lon: float | None) -> str | None:
    return None if lat is None or lon is None else f"{lat:.6f},{lon:.6f}"


def _num(v: Any) -> tuple[float | None, str | None]:
    if isblank(v):
        return None, None
    try:
        return float(str(v).strip()), None
    except ValueError:
        return None, f"cannot read '{str(v).strip()}'"


@dataclass
class Item:
    row_number: int
    raw: dict[str, Any]
    p: dict[str, Any]  # parsed fields; p["reasons"] are the parse-stage reasons
    id: int | None = None
    status: str = "ready"
    reason: str | None = None
    group: str = "unchanged"
    changes: list[dict[str, Any]] = field(default_factory=list)
    db_id: int | None = None
    note: str | None = None
    effective: tuple[float, float] | None = None
    rename: dict[str, Any] | None = None


@dataclass
class DbSite:
    id: int
    name: str
    kind: str
    active: bool
    lat: float | None
    lon: float | None
    details: dict[str, Any]
    lrd: str | None
    lrd_id: int | None

    @property
    def cls(self) -> str:
        return "dcs" if self.kind == "data_center" else "sites"


@dataclass
class ParsedSheet:
    kind: str
    sheet_name: str | None
    items: list[Item]
    found: bool = True


@dataclass
class SheetResult:
    kind: str
    sheet_name: str | None
    found: bool
    items: list[Item] = field(default_factory=list)
    missing: list[DbSite] = field(default_factory=list)
    in_use: set[int] = field(default_factory=set)
    renames: list[dict[str, Any]] = field(default_factory=list)
    block_reasons: list[str] = field(default_factory=list)
    hard: bool = False  # a block the admin cannot override

    @property
    def blocked(self) -> bool:
        return bool(self.block_reasons)

    @property
    def counts(self) -> dict[str, int]:
        c = Counter(i.group for i in self.items)
        return {"rows": len(self.items), "added": c["added"], "changed": c["changed"], "missing": len(self.missing),
                "problem": c["problem"], "unchanged": c["unchanged"]}

    def summary(self) -> dict[str, Any]:
        return {"kind": self.kind, "sheet_name": self.sheet_name, "found": self.found, "blocked": self.blocked,
                "block_reasons": self.block_reasons, "counts": self.counts,
                "needs_override": self.blocked and not self.hard}


# ---- reading the workbook -------------------------------------------------------------------------------------

def _find_sheet(wb: Any, kind: str, taken: set[str]) -> tuple[Any, int, list[str]] | None:
    need = set(SHEETS[kind][0])
    for ws in wb.worksheets:
        if ws.title in taken:
            continue
        for n, row in enumerate(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), 1):
            if need <= {hkey(c) for c in row if not blank(c)}:
                return ws, n, [" ".join(str(c).split()) if not blank(c) else "" for c in row]
    return None


def _parse_row(kind: str, headers: list[str], values: tuple[Any, ...]) -> Item:
    own = SHEETS[kind][3]
    f = {hkey(h): v for h, v in zip(headers, values) if h}
    item = Item(0, {h: _json(v) for h, v in zip(headers, values) if h and not blank(v)}, {})
    reasons: list[str] = []
    key = None if isblank(f.get("sitename" if kind == "sites" else "dcnamelink")) else str(f.get("sitename" if kind == "sites" else "dcnamelink")).strip()
    lat_src = f.get("latitudeindecimal") if kind == "sites" else next((f[k] for k in ("lat", "latitude") if k in f), None)
    lon_src = f.get("longitudeindecimal") if kind == "sites" else next((f[k] for k in ("long", "lon", "longitude") if k in f), None)
    lat, e1 = _num(lat_src)
    lon, e2 = _num(lon_src)
    coord_why = (f"latitude {e1}" if e1 else None) or (f"longitude {e2}" if e2 else None) \
        or (checks.check_coords(lat, lon) if kind == "sites" or lat is not None or lon is not None else None)
    if coord_why:
        lat = lon = None
    p = {"key": key, "lat": lat, "lon": lon,
         "details": {h: _json(v) for h, v in zip(headers, values) if h and hkey(h) not in own and not isblank(v)}}
    if not key:
        reasons.append(("site_name" if kind == "sites" else "DC Name / Link") + " is blank")
    if kind == "sites":
        if coord_why:
            reasons.append(coord_why)
        p["region"], why = checks.map_region(None if isblank(f.get("region")) else f.get("region"))
        if why:
            reasons.append(why)
    else:
        p["lrd"] = None if isblank(f.get("lrd")) else str(f["lrd"]).strip()
        p["own_coords_note"] = coord_why  # not a problem by itself: the DC falls back to its site
    p["reasons"] = reasons
    item.p = p
    return item


def parse_workbook(data: Any) -> dict[str, ParsedSheet]:
    """Both parts of the site master workbook. A part with no matching sheet comes back found=False. Never raises on a row."""
    wb = openpyxl.load_workbook(data, read_only=True, data_only=True)
    out: dict[str, ParsedSheet] = {}
    taken: set[str] = set()
    for kind in SHEETS:
        hit = _find_sheet(wb, kind, taken)
        if not hit:
            out[kind] = ParsedSheet(kind, None, [], found=False)
            continue
        ws, header_row, headers = hit
        taken.add(ws.title)
        items = []
        for n, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if all(blank(v) for v in values):
                continue
            try:
                it = _parse_row(kind, headers, values)
            except Exception as e:  # a bad row never stops the import
                it = Item(0, {h: _json(v) for h, v in zip(headers, values) if h and not blank(v)},
                          {"key": None, "details": {}, "reasons": [f"could not parse row: {type(e).__name__}: {e}"]})
            it.row_number = n
            items.append(it)
        out[kind] = ParsedSheet(kind, ws.title, items)
    return out


# ---- diff -----------------------------------------------------------------------------------------------------

def load_snapshot(db: Session) -> dict[str, DbSite]:
    rows = db.execute(text(
        "SELECT s.id, s.name, s.kind, s.active, ST_Y(s.geom) AS lat, ST_X(s.geom) AS lon, s.details, "
        "l.name AS lrd, s.lrd_site_id AS lrd_id FROM sites s LEFT JOIN sites l ON l.id = s.lrd_site_id")).mappings()
    return {r["name"]: DbSite(**{**r, "details": r["details"] or {}}) for r in rows}


def load_protected(db: Session) -> dict[int, set[str]]:
    prot: dict[int, set[str]] = {}
    for i, f in db.execute(text("SELECT entity_id, field FROM data_edits WHERE entity = 'site' AND reason NOT LIKE :r"),
                           {"r": UPLOAD_REASON + "%"}):
        prot.setdefault(i, set()).add(f)
    return prot


def is_protected(prot: dict[int, set[str]], site_id: int, fld: str) -> bool:
    fs = prot.get(site_id)
    if not fs:
        return False
    names = {fld}
    if fld.startswith("details."):
        names.add(fld[8:])
    if fld == "location":
        names |= LOCATION_FIELDS
    return bool(fs & names)


def _same(a: Any, b: Any) -> bool:
    return a == b or (a is not None and b is not None and str(a) == str(b))


def _new_fields(kind: str, p: dict[str, Any]) -> dict[str, Any]:
    f: dict[str, Any] = {"location": _fmt(p["lat"], p["lon"])}
    f.update({"region": p["region"]} if kind == "sites" else {"lrd": p["lrd"]})
    f.update({"details." + k: v for k, v in p["details"].items()})
    return f


def _old_fields(kind: str, row: DbSite) -> dict[str, Any]:
    details = dict(row.details)
    f: dict[str, Any] = {"location": _fmt(row.lat, row.lon)}
    if kind == "sites":
        rk = next((k for k in details if norm(k) == "region"), None)
        raw = details.pop(rk, None) if rk else None
        f["region"] = (checks.map_region(raw)[0] or raw) if raw is not None else None
    else:
        f["lrd"] = row.lrd
    f.update({"details." + k: v for k, v in details.items()})
    return f


def diff_fields(kind: str, p: dict[str, Any], row: DbSite, prot: dict[int, set[str]]) -> list[dict[str, Any]]:
    new, old = _new_fields(kind, p), _old_fields(kind, row)
    oldn = {norm(k): k for k in old}
    changes = []
    for k, v in new.items():
        ov = old.get(oldn.get(norm(k), ""))
        if not _same(ov, v):
            changes.append({"field": k, "old": ov, "new": v})
    newn = {norm(k) for k in new}
    changes += [{"field": k, "old": v, "new": None} for k, v in old.items() if norm(k) not in newn and v is not None]
    if not row.active:
        changes.append({"field": "active", "old": False, "new": True})
    for c in changes:
        c["protected"] = is_protected(prot, row.id, c["field"])
        if c["protected"]:
            c["note"] = "edited in dashboard, file differs"
    return changes


def _eval_sheet(kind: str, ps: ParsedSheet, snap: dict[str, DbSite], prot: dict[int, set[str]],
                site_coords: dict[str, tuple[float, float]], known: set[str]) -> SheetResult:
    res = SheetResult(kind, ps.sheet_name, ps.found, ps.items)
    expected = SHEETS[kind][1]
    if not ps.found:
        res.block_reasons, res.hard = [f"No worksheet has the expected headers: {expected}"], True
        return res
    label = "site_name" if kind == "sites" else "DC Name / Link"
    counts = Counter(i.p["key"] for i in ps.items if i.p["key"])
    dup = sorted(k for k, n in counts.items() if n > 1)
    for it in ps.items:
        p, reasons = it.p, list(it.p["reasons"])
        key = p["key"]
        if key and counts[key] > 1:
            reasons.append(f"duplicate {label} '{key}' in file (rows {', '.join(str(i.row_number) for i in ps.items if i.p['key'] == key)})")
        if kind == "dcs":
            lrd = p["lrd"]
            if not lrd:
                reasons.append("LRD is blank")
            elif lrd not in known:
                reasons.append(f"unknown site code {lrd}")
            if p["lat"] is not None:
                it.effective = (p["lat"], p["lon"])
            elif lrd in known:
                if lrd in site_coords:
                    it.effective, it.note = site_coords[lrd], f"location from site {lrd}"
                else:
                    reasons.append(f"no valid coordinates for the DC or its site {lrd}")
        row = snap.get(key) if key else None
        if row and row.cls != kind:
            reasons.append(f"name already used by a {row.kind} in the database")
        it.reason = "; ".join(reasons) or None
        it.status = "needs_review" if reasons else "ready"
        if reasons:
            it.group = "problem"
        elif row is None:
            it.group = "added"
        else:
            it.db_id, it.changes = row.id, diff_fields(kind, p, row, prot)
            it.group = "changed" if it.changes else "unchanged"
    res.missing = [r for r in snap.values() if r.cls == kind and r.active and r.name not in counts]
    # whole-file checks
    if not ps.items:
        res.block_reasons.append("sheet is empty")
        res.hard = True
    if dup:
        res.block_reasons.append(f"duplicate keys in file: {', '.join(dup[:5])}{' ...' if len(dup) > 5 else ''}")
        res.hard = True
    active = sum(1 for r in snap.values() if r.cls == kind and r.active)
    if ps.items and len(counts) * DROP_LIMIT[1] <= active * DROP_LIMIT[0]:
        res.block_reasons.append(f"file has {len(counts)} rows but the database has {active} active "
                                 f"({round(100 * (active - len(counts)) / active)}% fewer, limit is 10%)")
    if kind == "dcs":
        _possible_renames(res, site_coords)
    return res


def _possible_renames(res: SheetResult, site_coords: dict[str, tuple[float, float]]) -> None:
    added = [i for i in res.items if i.group == "added" and i.effective]
    taken: set[int] = set()
    for m in res.missing:
        loc = (m.lat, m.lon) if m.lat is not None else site_coords.get(m.lrd or "")
        if not loc:
            continue
        for it in added:
            if id(it) not in taken and it.p["lrd"] == m.lrd and _fmt(*it.effective) == _fmt(*loc):  # type: ignore[misc]
                taken.add(id(it))
                it.rename = {"missing_dc_id": m.id, "missing_dc_name": m.name}
                res.renames.append({"missing_dc_id": m.id, "missing_dc_name": m.name, "new_dc_name": it.p["key"],
                                    "new_row_number": it.row_number, "new_row_id": it.id, "lrd": m.lrd})
                break


def used_site_ids(db: Session, ids: list[int]) -> set[int]:
    if not ids:
        return set()
    sql = ("SELECT DISTINCT i FROM (SELECT site_a_id AS i FROM spans UNION ALL SELECT site_b_id FROM spans "
           "UNION ALL SELECT site_id FROM ports UNION ALL SELECT site_a_id FROM segments "
           "UNION ALL SELECT site_b_id FROM segments) u WHERE i = ANY(:ids)")
    return {r[0] for r in db.execute(text(sql), {"ids": ids})}


def evaluate(db: Session, parsed: dict[str, ParsedSheet]) -> dict[str, SheetResult]:
    snap, prot = load_snapshot(db), load_protected(db)
    coords = {n: (r.lat, r.lon) for n, r in snap.items() if r.cls == "sites" and r.lat is not None and r.lon is not None}
    known = {n for n, r in snap.items() if r.cls == "sites"}
    out = {"sites": _eval_sheet("sites", parsed["sites"], snap, prot, {}, set())}
    if not out["sites"].blocked or not out["sites"].hard:  # file sites count unless the sheet is blocked for good
        for it in out["sites"].items:
            if it.status == "ready":
                known.add(it.p["key"])
                if it.p["lat"] is not None:
                    coords[it.p["key"]] = (it.p["lat"], it.p["lon"])
    out["dcs"] = _eval_sheet("dcs", parsed["dcs"], snap, prot, coords, known)
    for r in out.values():
        r.in_use = used_site_ids(db, [m.id for m in r.missing])
    return out


# ---- staging --------------------------------------------------------------------------------------------------

def _payload(it: Item) -> dict[str, Any]:
    return {"key": it.p["key"], "group": it.group, "db_id": it.db_id, "changes": it.changes,
            "edited_in_dashboard": any(c["protected"] for c in it.changes), "location_note": it.note,
            "possible_rename": it.rename, "parsed": it.p}


def stage(db: Session, admin_id: int, filename: str, parsed: dict[str, ParsedSheet]) -> dict[str, Any]:
    """Store every row raw in upload_rows and the uploads row. Writes nothing else."""
    res = evaluate(db, parsed)
    summary = {"sheets": [res[k].summary() for k in SHEETS], "possible_renames": [], "applied": None}
    uid = db.execute(text("INSERT INTO uploads (user_id, filename, kind, status, summary) "
                          "VALUES (:u, :f, 'site_master', 'staged', CAST(:s AS jsonb)) RETURNING id"),
                     {"u": admin_id, "f": filename, "s": json.dumps(summary)}).scalar_one()
    params = []
    for kind, r in res.items():
        entity = SHEETS[kind][2]
        for it in r.items:
            params.append({"u": uid, "g": GROUP_COL[it.group], "e": entity, "p": json.dumps(_payload(it)),
                           "sh": r.sheet_name, "n": it.row_number, "st": it.status, "r": it.reason,
                           "raw": json.dumps(it.raw)})
        for m in r.missing:
            params.append({"u": uid, "g": "missing", "e": entity, "sh": r.sheet_name, "n": None, "st": None, "r": None,
                           "raw": None, "p": json.dumps({"key": m.name, "group": "missing", "db_id": m.id,
                                                         "in_use": m.id in r.in_use, "possible_rename": next(
                                                             (x for x in r.renames if x["missing_dc_id"] == m.id), None)})})
    if params:
        db.execute(text("INSERT INTO upload_rows (upload_id, row_group, entity, payload, sheet, row_number, status, reason, raw) "
                        "VALUES (:u, :g, :e, CAST(:p AS jsonb), :sh, :n, :st, :r, CAST(:raw AS jsonb))"), params)
    renames = res["dcs"].renames
    if renames:
        ids = dict(db.execute(text("SELECT row_number, id FROM upload_rows WHERE upload_id = :u AND entity = 'dc' "
                                   "AND row_group = 'new'"), {"u": uid}).all())
        for x in renames:
            x["new_row_id"] = ids.get(x["new_row_number"])
            db.execute(text("UPDATE upload_rows SET payload = jsonb_set(payload, '{possible_rename}', CAST(:x AS jsonb)) "
                            "WHERE id = :i"), {"x": json.dumps(x), "i": x["new_row_id"]})
            db.execute(text("UPDATE upload_rows SET payload = jsonb_set(payload, '{possible_rename}', CAST(:x AS jsonb)) "
                            "WHERE upload_id = :u AND row_group = 'missing' AND (payload->>'db_id')::bigint = :m"),
                       {"x": json.dumps(x), "u": uid, "m": x["missing_dc_id"]})
        summary["possible_renames"] = renames
        db.execute(text("UPDATE uploads SET summary = CAST(:s AS jsonb) WHERE id = :u"), {"s": json.dumps(summary), "u": uid})
    db.commit()
    return {"upload_id": uid, "filename": filename, **summary}


def load_staged(db: Session, upload_id: int, summary: dict[str, Any]) -> dict[str, ParsedSheet]:
    names = {s["kind"]: (s["sheet_name"], s["found"]) for s in summary["sheets"]}
    out = {k: ParsedSheet(k, names[k][0], [], names[k][1]) for k in SHEETS}
    ent = {v[2]: k for k, v in SHEETS.items()}
    for r in db.execute(text("SELECT id, entity, row_number, raw, payload FROM upload_rows "
                             "WHERE upload_id = :u AND row_group <> 'missing' ORDER BY id"), {"u": upload_id}).mappings():
        out[ent[r["entity"]]].items.append(Item(r["row_number"], r["raw"] or {}, r["payload"]["parsed"], id=r["id"]))
    return out


# ---- apply ----------------------------------------------------------------------------------------------------

def _apply_changes(kind: str, row: DbSite, it: Item, lrd_ids: dict[str, int], warnings: list[str]) -> tuple[dict[str, Any], list[tuple]]:
    """Final column values for a site and the data_edits rows (field, old, new) actually applied."""
    fin: dict[str, Any] = {"n": row.name, "lat": row.lat, "lon": row.lon, "d": dict(row.details), "l": row.lrd_id, "a": row.active}
    done: list[tuple] = []

    def setd(key: str, val: Any) -> None:
        for k in [k for k in fin["d"] if norm(k) == norm(key)]:
            del fin["d"][k]
        if val is not None:
            fin["d"][key] = val

    for c in it.changes:
        if c["protected"]:
            continue
        f, new = c["field"], c["new"]
        if f == "location":
            fin["lat"], fin["lon"] = it.p["lat"], it.p["lon"]
        elif f == "region":
            setd("region", new)
        elif f == "lrd":
            if new not in lrd_ids:
                warnings.append(f"{it.p['key']}: LRD site {new} was not found; link kept")
                continue
            fin["l"] = lrd_ids[new]
        elif f == "active":
            fin["a"] = True
        elif f == "name":
            fin["n"] = new
        else:
            setd(f[8:], new)
        done.append((f, c["old"], new))
    return fin, done


def apply_sheet(db: Session, admin_id: int, kind: str, r: SheetResult, reason: str, pairs: dict[int, Item],
                warnings: list[str], prot: dict[int, set[str]], snap: dict[str, DbSite]) -> dict[str, int]:
    out = dict.fromkeys(["added", "changed", "protected_kept", "renamed", "deactivated", "kept_active", "skipped"], 0)
    by_id = {s.id: s for s in snap.values()}
    # DC links resolve against the database now, so sites inserted earlier in this transaction count
    lrd_ids = dict(db.execute(text("SELECT name, id FROM sites WHERE kind <> 'data_center'")).all()) if kind == "dcs" else {}
    renamed_new = {id(i) for i in pairs.values()}
    inserts, updates, edits = [], [], []

    for it in r.items:
        if it.group == "added" and id(it) not in renamed_new:
            lid = lrd_ids.get(it.p["lrd"]) if kind == "dcs" else None
            if kind == "dcs" and lid is None:
                warnings.append(f"{it.p['key']}: LRD site {it.p['lrd']} is not in the database; DC skipped")
                out["skipped"] += 1
                continue
            d = dict(it.p["details"])
            if kind == "sites":
                d["region"] = it.p["region"]
            inserts.append({"n": it.p["key"], "k": "site" if kind == "sites" else "data_center", "lat": it.p["lat"],
                            "lon": it.p["lon"], "d": json.dumps(d), "l": lid})
        elif it.group == "changed":
            row = by_id[it.db_id]  # type: ignore[index]
            out["protected_kept"] += sum(c["protected"] for c in it.changes)
            fin, done = _apply_changes(kind, row, it, lrd_ids, warnings)
            if done:
                out["changed"] += 1
                updates.append({**fin, "d": json.dumps(fin["d"]), "i": row.id})
                edits += [(row.id, f, o, n) for f, o, n in done]
    # possible renames chosen as "rename": update the existing DC, keep its history
    for missing_id, it in pairs.items():
        row = by_id[missing_id]
        it.changes = diff_fields(kind, it.p, row, prot)
        if it.p["key"] != row.name:
            it.changes.append({"field": "name", "old": row.name, "new": it.p["key"], "protected": is_protected(prot, row.id, "name")})
        out["protected_kept"] += sum(c["protected"] for c in it.changes)
        fin, done = _apply_changes(kind, row, it, lrd_ids, warnings)
        out["renamed"] += 1
        if done:
            updates.append({**fin, "d": json.dumps(fin["d"]), "i": row.id})
            edits += [(row.id, f, o, n) for f, o, n in done]
    for m in r.missing:
        if m.id in pairs:
            continue
        if m.id in r.in_use:
            warnings.append(f"{m.name} is missing from the file but is used by a span, port or segment; kept active")
            out["kept_active"] += 1
        elif is_protected(prot, m.id, "active"):
            out["protected_kept"] += 1
        else:
            updates.append({"n": m.name, "lat": m.lat, "lon": m.lon, "d": json.dumps(m.details), "l": m.lrd_id, "a": False, "i": m.id})
            edits.append((m.id, "active", True, False))
            out["deactivated"] += 1
    if inserts:
        db.execute(text("INSERT INTO sites (name, kind, geom, details, lrd_site_id) VALUES (:n, :k, "
                        "ST_SetSRID(ST_MakePoint(CAST(:lon AS double precision), CAST(:lat AS double precision)), 4326), "
                        "CAST(:d AS jsonb), :l)"), inserts)
        out["added"] = len(inserts)
    if updates:
        db.execute(text("UPDATE sites SET name = :n, geom = ST_SetSRID(ST_MakePoint(CAST(:lon AS double precision), "
                        "CAST(:lat AS double precision)), 4326), details = CAST(:d AS jsonb), lrd_site_id = :l, active = :a "
                        "WHERE id = :i"), updates)
    if edits:
        db.execute(text("INSERT INTO data_edits (entity, entity_id, field, old_value, new_value, reason, user_id) "
                        "VALUES ('site', :i, :f, :o, :n, :r, :u)"),
                   [{"i": i, "f": f, "o": None if o is None else str(o).lower() if isinstance(o, bool) else str(o),
                     "n": None if n is None else str(n).lower() if isinstance(n, bool) else str(n), "r": reason, "u": admin_id}
                    for i, f, o, n in edits])
    return out


class ApplyError(Exception):
    def __init__(self, status: int, message: str):
        self.status, self.message = status, message


def apply_upload(db: Session, admin_id: int, upload_id: int, override: dict[str, bool], reason: str | None,
                 renames: list[dict[str, Any]]) -> dict[str, Any]:
    up = db.execute(text("SELECT id, filename, status, summary FROM uploads WHERE id = :i AND kind = 'site_master'"),
                    {"i": upload_id}).mappings().first()
    if not up:
        raise ApplyError(404, "Upload not found")
    if up["status"] != "staged":
        raise ApplyError(409, "This upload was already applied")
    if any(override.values()) and not (reason or "").strip():
        raise ApplyError(422, "A reason is required to override the row-count check")
    res = evaluate(db, load_staged(db, upload_id, up["summary"]))
    todo = [k for k in SHEETS if res[k].found and (not res[k].blocked or (not res[k].hard and override.get(k)))]
    if not todo:
        raise ApplyError(409, "Nothing can be applied: " + "; ".join(
            f"{k}: {', '.join(r.block_reasons)}" for k, r in res.items() if r.blocked))
    chosen = {x["missing_dc_id"]: x for x in renames if x["decision"] == "rename"}
    possible = {x["missing_dc_id"]: x for x in res["dcs"].renames}
    for m, x in chosen.items():
        if m not in possible or possible[m]["new_row_id"] != x["new_row_id"]:
            raise ApplyError(422, f"Missing DC {m} and row {x['new_row_id']} are not a possible rename in this upload")
    decided = {x["missing_dc_id"] for x in renames}
    undecided = [x for m, x in possible.items() if m not in decided] if "dcs" in todo else []
    if undecided:
        raise ApplyError(422, "Choose rename or separate for each possible rename: " + "; ".join(
            f"{x['missing_dc_name']} -> {x['new_dc_name']} (missing_dc_id {x['missing_dc_id']}, new_row_id {x['new_row_id']})"
            for x in undecided))
    why = f"{UPLOAD_REASON} {up['filename']} #{upload_id}"
    snap, prot = load_snapshot(db), load_protected(db)
    warnings: list[str] = []
    summary: dict[str, Any] = {}
    with db.begin_nested():  # all or nothing
        for k in todo:
            pairs = {}
            if k == "dcs":
                pairs = {m: next(i for i in res["dcs"].items if i.id == x["new_row_id"]) for m, x in chosen.items()}
            summary[k] = apply_sheet(db, admin_id, k, res[k], why, pairs, warnings, prot, snap)
        sheets = {k: {**summary[k], "applied": True} if k in summary else {"applied": False, "block_reasons": res[k].block_reasons}
                  for k in SHEETS}
        applied = {"by": admin_id, "at": datetime.now(timezone.utc).isoformat(), "sheets": sheets, "warnings": warnings,
                   "override_drop": {k: bool(override.get(k)) for k in SHEETS}, "override_reason": reason}
        for k, on in override.items():
            if on and k in summary:
                db.execute(text("INSERT INTO data_edits (entity, entity_id, field, old_value, new_value, reason, user_id) "
                                "VALUES ('upload', :i, 'override_drop', NULL, :k, :r, :u)"),
                           {"i": upload_id, "k": k, "r": reason.strip(), "u": admin_id})  # type: ignore[union-attr]
        s = dict(up["summary"], applied=applied)
        db.execute(text("UPDATE uploads SET status = 'confirmed', confirmed_at = now(), summary = CAST(:s AS jsonb) WHERE id = :i"),
                   {"s": json.dumps(s), "i": upload_id})
    db.commit()
    invalidate_routes()
    return {"upload_id": upload_id, "sheets": sheets, "warnings": warnings}
