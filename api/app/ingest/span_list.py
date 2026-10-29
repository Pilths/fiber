"""Span list upload: create and update spans (site pairs) from one sheet with Link and Distance headers.

Staging writes only `uploads` and `upload_rows`. Apply recomputes the diff from the staged rows against the
current database (like the site master), so a stale preview cannot overwrite newer data.
"""
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

import openpyxl
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ingest.checks import blank, split_nodes
from app.ingest.excel import _json
from app.ingest.reservation_upload import REASON as RESERVATION_REASON
from app.ingest.site_master import ApplyError, hkey
from app.routing.graph import invalidate as invalidate_routes

UPLOAD_REASON = "span list upload"
ALIASES = {"link": "link", "links": "link", "sitea": "site_a", "section": "section", "segment": "section",
           "distance": "distance", "distancekm": "distance", "km": "distance", "route": "route", "type": "type", "routetype": "type"}
EXPECTED = "Link, Distance"
GROUP_COL = {"new": "new", "changed": "changed", "unchanged": "unchanged", "problem": "needs_review"}
TYPES = {"trunk": "Trunk", "metro": "Metro"}
FIELDS = ("distance_km", "route_type", "route", "section")  # span columns a row may change
MAX_DISTANCE = Decimal("9999999.999")  # numeric(10,3)


@dataclass
class Item:
    row_number: int
    raw: dict[str, Any]
    p: dict[str, Any]
    id: int | None = None
    status: str = "ready"
    reason: str | None = None
    group: str = "unchanged"
    changes: list[dict[str, Any]] = field(default_factory=list)
    db_id: int | None = None
    ids: tuple[int, int] | None = None  # site ids of the pair, set when both sites exist


@dataclass
class Parsed:
    sheet_name: str | None
    items: list[Item]
    found: bool = True


def _cell(v: Any) -> str | None:
    return None if blank(v) else " ".join(str(v).split())


def _distance(v: Any) -> tuple[str | None, str | None]:
    """(distance text with 3 decimals, reason). Digits with '.' as the decimal point; blank is allowed."""
    if blank(v):
        return None, None
    s = format(Decimal(repr(v)), "f") if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v).strip()
    try:
        d = Decimal(s).quantize(Decimal("0.001"), ROUND_HALF_UP) if re.fullmatch(r"\d+(\.\d+)?", s) else None
    except InvalidOperation:
        d = None
    if d is None or d > MAX_DISTANCE:
        return None, "distance is not a number"
    return str(d), None


def _parse_row(headers: list[str], values: tuple[Any, ...]) -> Item:
    f: dict[str, Any] = {}
    for h, v in zip(headers, values):
        if h and ALIASES.get(hkey(h)):
            f.setdefault(ALIASES[hkey(h)], v)
    item = Item(0, {h: _json(v) for h, v in zip(headers, values) if h and not blank(v)}, {})
    reasons: list[str] = []
    link = _cell(f.get("link"))
    a = b = None
    if link is None:
        reasons.append("Link is blank")
    else:
        pair = None if re.search(r"\s", str(f["link"]).strip()) else split_nodes(link)
        if pair is None:
            reasons.append("Link must be two different site names joined by a dash, like AAAA-BBBB, with no spaces")
        else:
            a, b = pair[0], pair[1]
    dist, why = _distance(f.get("distance"))
    if why:
        reasons.append(why)
    typ = _cell(f.get("type"))
    if typ is not None and typ.casefold() not in TYPES:
        reasons.append("type must be Trunk or Metro")
    item.p = {"link": link, "a": a, "b": b, "distance": dist, "route": _cell(f.get("route")), "section": _cell(f.get("section")),
              "type": TYPES.get((typ or "").casefold()), "reasons": reasons}
    return item


def _find_sheet(wb: Any) -> tuple[Any, int, list[str]] | None:
    """First sheet with a row holding the Link header (any column) and a Distance header."""
    for ws in wb.worksheets:
        for n, row in enumerate(ws.iter_rows(values_only=True), 1):
            headers = [" ".join(str(c).split()) if not blank(c) else "" for c in row]
            fields = {ALIASES.get(hkey(h)) for h in headers if h}
            if {"link", "distance"} <= fields:
                return ws, n, headers
    return None


def parse_workbook(data: Any) -> Parsed:
    """Never raises on a row."""
    wb = openpyxl.load_workbook(data, read_only=True, data_only=True)
    hit = _find_sheet(wb)
    if not hit:
        return Parsed(None, [], found=False)
    ws, header_row, headers = hit
    items = []
    for n, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        if all(blank(v) for v in values):
            continue
        try:
            it = _parse_row(headers, values)
        except Exception as e:  # a bad row never stops the import
            it = Item(0, {h: _json(v) for h, v in zip(headers, values) if h and not blank(v)},
                      {"link": None, "a": None, "b": None, "distance": None, "route": None, "section": None, "type": None,
                       "reasons": [f"could not parse row: {type(e).__name__}: {e}"]})
        it.row_number = n
        items.append(it)
    return Parsed(ws.title, items)


# ---- diff -----------------------------------------------------------------------------------------------------

def load_protected(db: Session) -> dict[int, set[str]]:
    prot: dict[int, set[str]] = {}
    for i, f in db.execute(text("SELECT entity_id, field FROM data_edits WHERE entity = 'span' AND reason NOT LIKE :a AND reason NOT LIKE :b"),
                           {"a": UPLOAD_REASON + "%", "b": RESERVATION_REASON + "%"}):
        prot.setdefault(i, set()).add(f)
    return prot


def _old(row: Any) -> dict[str, str | None]:
    return {k: None if row[k] is None else str(row[k]) for k in FIELDS}


def evaluate(db: Session, parsed: Parsed) -> list[Item]:
    """Set status, reason, group, changes and db_id on every item."""
    names = {n for it in parsed.items for n in (it.p["a"], it.p["b"]) if n}
    sites = {r.name: r for r in db.execute(text("SELECT id, name, active FROM sites WHERE name = ANY(:n)"), {"n": sorted(names)})}
    spans = {frozenset((r["site_a_id"], r["site_b_id"])): r for r in db.execute(text(
        "SELECT id, site_a_id, site_b_id, distance_km, route_type, route, section FROM spans")).mappings()}
    prot = load_protected(db)
    pairs = Counter(frozenset((it.p["a"], it.p["b"])) for it in parsed.items if it.p["a"])
    for it in parsed.items:
        p, reasons = it.p, list(it.p["reasons"])
        if p["a"]:
            key = frozenset((p["a"], p["b"]))
            if pairs[key] > 1:
                rows_ = ", ".join(str(i.row_number) for i in parsed.items if i.p["a"] and frozenset((i.p["a"], i.p["b"])) == key)
                reasons.append(f"duplicate link in file (rows {rows_})")
            for n in (p["a"], p["b"]):
                if n not in sites:
                    reasons.append(f"site {n} is not in the dashboard")
                elif not sites[n].active:
                    reasons.append(f"site {n} is not active in the dashboard")
        it.reason = "; ".join(reasons) or None
        it.status = "needs_review" if reasons else "ready"
        if reasons:
            it.group = "problem"
            continue
        ia, ib = sites[p["a"]].id, sites[p["b"]].id
        it.ids = (ia, ib)
        row = spans.get(frozenset((ia, ib)))
        if row is None:
            it.group = "new"
            continue
        it.db_id, old = row["id"], _old(row)
        new = {"distance_km": p["distance"], "route_type": p["type"], "route": p["route"], "section": p["section"]}
        for f in FIELDS:  # a blank file cell never clears a value
            if new[f] is not None and new[f] != old[f]:
                prot_f = f in prot.get(row["id"], ())
                it.changes.append({"field": f, "old": old[f], "new": new[f], "protected": prot_f,
                                   "note": "edited in dashboard, file differs" if prot_f else None})
        it.group = "changed" if it.changes else "unchanged"
    return parsed.items


def block_reason(parsed: Parsed) -> str | None:
    if not parsed.found:
        return f"No worksheet has the expected headers: {EXPECTED}"
    return None if parsed.items else "sheet is empty"


def _counts(items: list[Item]) -> dict[str, int]:
    c = Counter(i.group for i in items)
    return {"rows": len(items), "new": c["new"], "changed": c["changed"], "unchanged": c["unchanged"], "problem": c["problem"]}


# ---- staging --------------------------------------------------------------------------------------------------

def stage(db: Session, admin_id: int, filename: str, parsed: Parsed) -> dict[str, Any]:
    """Store every row raw in upload_rows and the uploads row. Writes nothing else."""
    items = evaluate(db, parsed)
    summary = {"sheet_name": parsed.sheet_name, "found": parsed.found, "block_reason": block_reason(parsed),
               "counts": _counts(items), "applied": None}
    uid = db.execute(text("INSERT INTO uploads (user_id, filename, kind, status, summary) "
                          "VALUES (:u, :f, 'span_list', 'staged', CAST(:s AS jsonb)) RETURNING id"),
                     {"u": admin_id, "f": filename, "s": json.dumps(summary)}).scalar_one()
    if items:
        db.execute(text("INSERT INTO upload_rows (upload_id, row_group, entity, payload, sheet, row_number, status, reason, raw) "
                        "VALUES (:u, :g, 'span', CAST(:p AS jsonb), :sh, :n, :st, :r, CAST(:raw AS jsonb))"),
                   [{"u": uid, "g": GROUP_COL[it.group], "sh": parsed.sheet_name, "n": it.row_number, "st": it.status,
                     "r": it.reason, "raw": json.dumps(it.raw),
                     "p": json.dumps({"group": it.group, "db_id": it.db_id, "changes": it.changes,
                                      "edited_in_dashboard": any(c["protected"] for c in it.changes), "parsed": it.p})}
                    for it in items])
    db.commit()
    return {"upload_id": uid, "filename": filename, **summary}


def load_staged(db: Session, upload_id: int, summary: dict[str, Any]) -> Parsed:
    items = [Item(r["row_number"], r["raw"] or {}, r["payload"]["parsed"], id=r["id"]) for r in db.execute(text(
        "SELECT id, row_number, raw, payload FROM upload_rows WHERE upload_id = :u ORDER BY id"), {"u": upload_id}).mappings()]
    return Parsed(summary["sheet_name"], items, summary["found"])


# ---- apply ----------------------------------------------------------------------------------------------------

def apply_upload(db: Session, admin_id: int, upload_id: int) -> dict[str, Any]:
    up = db.execute(text("SELECT id, filename, status, summary FROM uploads WHERE id = :i AND kind = 'span_list' FOR UPDATE"),
                    {"i": upload_id}).mappings().first()
    if not up:
        raise ApplyError(404, "Upload not found")
    if up["status"] != "staged":
        raise ApplyError(409, "This upload was already applied")
    parsed = load_staged(db, upload_id, up["summary"])
    if why := block_reason(parsed):
        raise ApplyError(409, f"Nothing can be applied: {why}")
    items = evaluate(db, parsed)
    why = f"{UPLOAD_REASON} {up['filename']} #{upload_id}"
    out = {"created": 0, "changed": 0, "protected_kept": 0, "skipped": 0}
    warnings: list[str] = []
    edits: list[dict[str, Any]] = []
    with db.begin_nested():  # all or nothing
        for it in items:
            if it.group == "problem":
                out["skipped"] += 1
            elif it.group == "new":
                p = it.p
                db.execute(text("INSERT INTO spans (site_a_id, site_b_id, route_type, distance_km, route, section) "
                                "VALUES (:a, :b, :t, CAST(:d AS numeric), :r, :s)"),
                           {"a": it.ids[0], "b": it.ids[1], "t": p["type"], "d": p["distance"], "r": p["route"], "s": p["section"]})  # type: ignore[index]
                out["created"] += 1
            elif it.group == "changed":
                out["protected_kept"] += sum(c["protected"] for c in it.changes)
                todo = [c for c in it.changes if not c["protected"]]
                if not todo:
                    continue
                for c in todo:  # c["field"] is one of FIELDS
                    db.execute(text(f"UPDATE spans SET {c['field']} = CAST(:v AS {'numeric' if c['field'] == 'distance_km' else 'text'}) WHERE id = :i"),
                               {"v": c["new"], "i": it.db_id})
                    edits.append({"i": it.db_id, "f": c["field"], "o": c["old"], "n": c["new"], "r": why, "u": admin_id})
                out["changed"] += 1
        if edits:
            db.execute(text("INSERT INTO data_edits (entity, entity_id, field, old_value, new_value, reason, user_id) "
                            "VALUES ('span', :i, :f, :o, :n, :r, :u)"), edits)
        applied = {"by": admin_id, "at": datetime.now(UTC).isoformat(), **out, "warnings": warnings}
        db.execute(text("UPDATE uploads SET status = 'confirmed', confirmed_at = now(), summary = CAST(:s AS jsonb) WHERE id = :i"),
                   {"s": json.dumps(dict(up["summary"], applied=applied)), "i": upload_id})
    db.commit()
    invalidate_routes()
    return {"upload_id": upload_id, **out, "warnings": warnings}
