"""Parse the sample workbook (sheets Reservation, Site Location, DC List) into row results. Reads only; never raises on a row."""
import itertools
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel

from app.ingest import checks
from app.ingest.checks import blank, norm

SITE_HEADER_ROW = 4


@dataclass
class RowResult:
    sheet: str
    row_number: int
    raw: dict[str, Any]
    parsed: dict[str, Any] = field(default_factory=dict)
    status: str = "ready"
    reason: str | None = None


def _json(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v if v is None or isinstance(v, (str, int, float, bool)) else str(v)


def _float(v: Any) -> float | None:
    return None if blank(v) else float(str(v).strip())


def _date(v: Any) -> date | None:
    if blank(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return from_excel(v).date()
    return date.fromisoformat(str(v).strip())


def _text(v: Any) -> str | None:
    return None if blank(v) else str(v).strip()


def _read(ws: Any, header_row: int) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    """Header texts and (row number, values), with merged cells filled from their top-left cell."""
    merged: dict[tuple[int, int], Any] = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for rc in itertools.product(range(rng.min_row, rng.max_row + 1), range(rng.min_col, rng.max_col + 1)):
            merged[rc] = v
    rows = [(r[0].row, [merged.get((c.row, c.column), c.value) for c in r]) for r in ws.iter_rows(min_row=header_row)]
    headers = [" ".join(str(v).split()) if not blank(v) else "" for v in rows[0][1]] if rows else []
    return headers, rows[1:]


def _sheet(wb: Any, name: str) -> Any:
    return next((ws for ws in wb.worksheets if norm(ws.title) == norm(name)), None)


def _header_row(ws: Any, anchor: str) -> int | None:
    """Number of the first row, top to bottom, that holds the anchor header in any column; None when no row does."""
    return next((r[0].row for r in ws.iter_rows() if any(not blank(c.value) and norm(c.value) == anchor for c in r)), None)


def _row(sheet: str, n: int, headers: list[str], values: list[Any]) -> tuple[RowResult, dict[str, Any]]:
    raw = {h: _json(v) for h, v in zip(headers, values) if h}
    return RowResult(sheet, n, raw), {norm(h): v for h, v in zip(headers, values) if h}


def _finish(r: RowResult, reasons: list[str], reject: bool = False) -> None:
    r.reason = "; ".join(reasons) or None
    r.status = "rejected" if reject else "needs_review" if reasons else "ready"


def _guard(r: RowResult, fn: Any) -> None:
    try:
        fn()
    except Exception as e:  # a bad row never stops the import
        r.parsed, r.status, r.reason = {}, "needs_review", f"could not parse row: {type(e).__name__}: {e}"


def _reservations(ws: Any, known: set[str]) -> list[RowResult]:
    headers, rows = _read(ws, _header_row(ws, "nodes") or 1)  # Nodes is the anchor header; other columns sit beside it
    out: list[RowResult] = []
    label: str | None = None
    for n, values in rows:
        filled = [v for v in values if not blank(v)]
        if not filled:
            continue
        r, f = _row("Reservation", n, headers, values)
        if len(filled) == 1 and isinstance(filled[0], str) and blank(f.get("nodes")):
            label = filled[0].strip()
            continue
        _guard(r, lambda: _reservation(r, f, label, known))
        out.append(r)
    return out


def _reservation(r: RowResult, f: dict[str, Any], label: str | None, known: set[str]) -> None:
    reasons: list[str] = []
    no_customer = checks.check_requestor(f.get("fiber core requestor"))
    region, why = checks.map_region(f.get("region"))
    nodes = checks.split_nodes(f.get("nodes"))
    p = r.parsed = {"group_label": label, "region": region}
    if why:
        reasons.append(why)
    if nodes:
        p["site_a"], p["site_b"], p["additional"] = nodes
        if why := checks.check_known_sites(nodes[:2], known):
            reasons.append(why)
    else:
        reasons.append("Nodes is not in the form AABB-BBCC")
    # every row is a count-only reservation; Core assignment and Port info stay raw (ports are added by hand later)
    count = f.get("no of core requested")
    p["unassigned_cores"] = 0
    if blank(count):
        reasons.append("No of core requested is blank")
    else:
        try:
            n = float(str(count).strip())
            if n < 1 or n != int(n):
                raise ValueError
            p["unassigned_cores"] = int(n)
        except ValueError:
            reasons.append(f"No of core requested '{count}' is not a whole number of 1 or more")
    p["customer"] = _text(f.get("fiber core requestor"))
    p["contractor"], p["remark"] = _text(f.get("contractor")), _text(f.get("remark"))
    p["port_info_raw"] = _text(f.get("port info"))
    p["monitoring_code"] = _text(f.get("fiber core monitoring"))
    fm = _text(f.get("fiber montiored y/n/wip", f.get("fiber monitored y/n/wip")))
    p["fiber_monitored"] = fm and fm.upper()
    if p["fiber_monitored"] not in (None, "Y", "N", "WIP"):
        reasons.append(f"fiber monitored '{fm}' is not Y, N or WIP")
    for key, col, conv in [("distance_km", "distance(km)", _float), ("survey_cores", "core secured", _float),
                           ("cv_ready_date", "expected cv ready", _date), ("planned_date", "fiber core request date", _date)]:
        try:
            p[key] = conv(f.get(col))
        except ValueError:
            p[key] = None
            reasons.append(f"cannot read {col} '{f.get(col)}'")
    if p["survey_cores"] is not None:
        p["survey_cores"] = int(p["survey_cores"])
    # optional columns (not in the sample workbook): a test result for the span, and a core fault mark
    for key, col, conv in [("loss_db", "loss (db)", _float), ("tested_on", "tested on", _date)]:
        try:
            p[key] = conv(f.get(col, f.get(col.replace(" (", "("))))
        except ValueError:
            p[key] = None
            reasons.append(f"cannot read {col} '{f.get(col)}'")
    cs = _text(f.get("core status"))
    p["core_status"] = cs and cs.capitalize()
    if cs and p["core_status"] != "Faulty":
        reasons.append(f"core status '{cs}' is not Faulty")
    if no_customer:
        reasons.insert(0, no_customer)
    _finish(r, reasons, reject=bool(no_customer))


def _sites(ws: Any, header_row: int, sheet: str, seen: set[str]) -> list[RowResult]:
    headers, rows = _read(ws, header_row)
    out: list[RowResult] = []
    dc = sheet == "DC List"
    name_col, lat_col, lon_col = ("dc name / link", "lat", "long") if dc else ("site_name", "latitude_in_decimal", "longitude_in_decimal")
    own = {name_col, lat_col, lon_col, "lrd"} if dc else {name_col, lat_col, lon_col}
    for n, values in rows:
        r, f = _row(sheet, n, headers, values)
        if blank(f.get(name_col)):
            continue
        name = str(f[name_col]).strip()

        def parse() -> None:
            reasons = []
            lat, lon = _float(f.get(lat_col)), _float(f.get(lon_col))
            if why := checks.check_coords(lat, lon):
                reasons.append(why)
            if name in seen:
                reasons.append("duplicate site name")
            seen.add(name)
            details = {h: _json(v) for h, v in zip(headers, values) if h and norm(h) not in own and not blank(v)}
            r.parsed = {"name": name, "kind": "data_center" if dc else "site", "lat": lat, "lon": lon, "details": details}
            if dc:
                r.parsed["lrd"] = _text(f.get("lrd"))
            _finish(r, reasons)

        _guard(r, parse)
        out.append(r)
    return out


def parse_reservation_workbook(path: Any) -> tuple[list[RowResult], list[str]]:
    """Reservation sheet only (the site master comes from its own upload). Returns rows and the names of the sheets left out.

    Site codes are not judged here: the upload checks them against the database.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _sheet(wb, "Reservation") or next((w for w in wb.worksheets if _header_row(w, "nodes")), None)  # else any sheet with Nodes
    if ws is None:
        raise ValueError("no Reservation sheet found, and no sheet has a Nodes column")
    return _reservations(ws, set()), [w.title for w in wb.worksheets if w is not ws]


def parse_workbook(path: Any, known_sites: set[str] | None = None) -> list[RowResult]:
    """Rows of Site Location, DC List and Reservation (in that order). Monitoring is skipped.

    Site codes are judged against known_sites plus every site name listed in the workbook; with neither, no judgement.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[RowResult] = []
    seen: set[str] = set()
    for sheet, header_row in [("Site Location", SITE_HEADER_ROW), ("DC List", 1)]:
        if ws := _sheet(wb, sheet):
            out += _sites(ws, header_row, sheet, seen)
    known = set(known_sites or ()) | {r.parsed["name"] for r in out if "name" in r.parsed}
    for r in out:
        if r.sheet == "DC List" and r.parsed.get("lrd") and (why := checks.check_known_sites([r.parsed["lrd"]], known)):
            r.status, r.reason = "needs_review", "; ".join(filter(None, [r.reason, why]))
    if ws := _sheet(wb, "Reservation"):
        out += _reservations(ws, known)
    return out
